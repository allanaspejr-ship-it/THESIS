from __future__ import annotations

import json
import shutil
from pathlib import Path

import pandas as pd


APPENDIX_DIR = Path(__file__).resolve().parent
THESIS_DIR = APPENDIX_DIR.parent
TABLES_DIR = APPENDIX_DIR / "tables"
FIGURES_DIR = APPENDIX_DIR / "figures"
SCREENSHOTS_DIR = APPENDIX_DIR / "screenshots"
LATEX_DIR = APPENDIX_DIR / "latex"

PLANTS = ["Agus 1", "Agus 2", "Agus 4", "Agus 5", "Agus 6", "Agus 7"]
PLANT_KEYS = ["agus1", "agus2", "agus4", "agus5", "agus6", "agus7"]

GENERATED: list[Path] = []
WARNINGS: list[str] = []


def warn(message: str) -> None:
    WARNINGS.append(message)
    print(f"WARNING: {message}")


def ensure_dirs() -> None:
    for path in [
        TABLES_DIR,
        FIGURES_DIR,
        SCREENSHOTS_DIR,
        SCREENSHOTS_DIR / "gui_reference",
        FIGURES_DIR / "appendix_I_testing_plots",
        TABLES_DIR / "picture_tables",
        LATEX_DIR,
    ]:
        path.mkdir(parents=True, exist_ok=True)


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(THESIS_DIR)).replace("\\", "/")
    except ValueError:
        return str(path).replace("\\", "/")


def read_excel_safe(path: Path, nrows: int | None = None) -> pd.DataFrame:
    if not path.exists():
        warn(f"Missing source file: {rel(path)}")
        return pd.DataFrame()
    try:
        return pd.read_excel(path, nrows=nrows)
    except Exception as exc:
        warn(f"Could not read {rel(path)}: {exc}")
        return pd.DataFrame()


def read_json_safe(path: Path) -> dict:
    if not path.exists():
        warn(f"Missing source file: {rel(path)}")
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        warn(f"Could not read {rel(path)}: {exc}")
        return {}


def format_workbook(path: Path) -> None:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        warn(f"openpyxl unavailable for workbook formatting: {exc}")
        return

    try:
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row and ws.max_column:
                ws.auto_filter.ref = ws.dimensions
            for column_cells in ws.columns:
                max_len = 0
                letter = column_cells[0].column_letter
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, min(len(value), 80))
                ws.column_dimensions[letter].width = max(12, min(max_len + 2, 45))
        wb.save(path)
    except Exception as exc:
        warn(f"Could not format workbook {rel(path)}: {exc}")


def save_table(df: pd.DataFrame, stem: str, csv: bool = True, xlsx: bool = True) -> None:
    if xlsx:
        xlsx_path = TABLES_DIR / f"{stem}.xlsx"
        df.to_excel(xlsx_path, index=False)
        format_workbook(xlsx_path)
        GENERATED.append(xlsx_path)
    if csv:
        csv_path = TABLES_DIR / f"{stem}.csv"
        df.to_csv(csv_path, index=False)
        GENERATED.append(csv_path)


def save_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)
    format_workbook(path)
    GENERATED.append(path)


def make_table_png(df: pd.DataFrame, path: Path, max_rows: int = 24) -> None:
    if df.empty:
        warn(f"Skipped screenshot placeholder because table is empty: {rel(path)}")
        return
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception as exc:
        warn(f"matplotlib unavailable; could not create {rel(path)}: {exc}")
        return

    sample = df.head(max_rows).copy()
    sample = sample.fillna("")
    display_df = sample.copy()
    for col in display_df.columns:
        display_df[col] = display_df[col].astype(str).str.slice(0, 28)

    width = max(12, min(26, len(display_df.columns) * 1.35))
    height = max(4, min(18, (len(display_df) + 1) * 0.42))
    fig, ax = plt.subplots(figsize=(width, height), dpi=180)
    ax.axis("off")
    table = ax.table(
        cellText=display_df.values,
        colLabels=display_df.columns,
        loc="center",
        cellLoc="center",
    )
    table.auto_set_font_size(False)
    table.set_fontsize(7)
    table.scale(1, 1.25)
    for (row, _col), cell in table.get_celld().items():
        cell.set_edgecolor("#B8B8B8")
        if row == 0:
            cell.set_facecolor("#D9EAF7")
            cell.set_text_props(weight="bold")
        else:
            cell.set_facecolor("#FFFFFF" if row % 2 else "#F7F7F7")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    GENERATED.append(path)


def appendix_a() -> tuple[pd.DataFrame, pd.DataFrame]:
    rules = pd.DataFrame(
        [
            ["Flattening of three-row Excel header", "Merged multi-row Excel headers into a single machine-readable row.", "Create consistent columns for automated cleaning.", "scripts/cell1_clean_data.py; outputs/cleaned_data/cleaned_hourly_data.xlsx"],
            ["Standardization of plant, unit, generation, outage, gate, elevation, rainfall, and outflow column names", "Converted source labels into normalized snake_case field names.", "Reduce ambiguity across plant and hydrologic variables.", "outputs/cleaned_data/cleaned_hourly_data.xlsx"],
            ["Date and time conversion into hourly datetime", "Parsed date and time fields into hourly chronological records.", "Support time-series alignment, lag creation, and resampling.", "outputs/cleaned_data/cleaned_hourly_data.xlsx"],
            ["Near-zero generation treatment using 0.05 MW threshold", "Treated generation values near zero as operational zero values.", "Avoid numerical noise from being interpreted as true generation.", "scripts/cell1_clean_data.py"],
            ["Short missing-gap interpolation up to 2 hours", "Interpolated short missing gaps with a maximum length of two hours.", "Fill isolated missing values without inventing long outages.", "scripts/cell1_clean_data.py"],
            ["Short zero-run interpolation up to 2 hours", "Interpolated zero runs up to two consecutive hours.", "Smooth brief instrumentation dropouts.", "scripts/cell1_clean_data.py"],
            ["Preservation of long zero runs of at least 3 hours", "Kept zero runs of three hours or longer unchanged.", "Preserve likely outage or non-generation events.", "scripts/cell1_clean_data.py"],
            ["Spike smoothing using local-neighbor rule with factor 3.5", "Smoothed local spikes exceeding the neighbor-based threshold.", "Reduce isolated measurement spikes while retaining trend behavior.", "scripts/cell1_clean_data.py"],
            ["KNN imputation with 5 neighbors", "Used five-neighbor imputation for remaining missing values.", "Recover incomplete operational and hydrologic records.", "scripts/cell1_clean_data.py"],
            ["Hourly resampling", "Resampled cleaned records to hourly frequency.", "Ensure uniform time steps for forecasting.", "outputs/cleaned_data/cleaned_hourly_data.xlsx"],
            ["Recalculation of plant total generation from unit-level generation", "Recomputed plant totals from cleaned unit-level generation.", "Keep total generation consistent with unit records.", "outputs/cleaned_data/cleaned_hourly_data.xlsx"],
            ["Conversion of outage columns into binary unit status indicators", "Encoded outage/unit availability as binary status fields.", "Provide model-ready unit availability inputs.", "outputs/outages_planning/Planned_Outages_Input.xlsx"],
        ],
        columns=["Cleaning Step", "Rule Applied", "Purpose", "Source/Output File"],
    )
    save_table(rules, "appendix_A_data_cleaning_rules")

    cleaned_xlsx = THESIS_DIR / "outputs/cleaned_data/cleaned_hourly_data.xlsx"
    cleaned_parquet = THESIS_DIR / "outputs/cleaned_data/cleaned_hourly_data.parquet"
    outage_path = THESIS_DIR / "outputs/outages_planning/Planned_Outages_Input.xlsx"
    metadata = read_json_safe(THESIS_DIR / "outputs/cleaned_data/cell1_metadata.json")

    rows = metadata.get("rows")
    columns = metadata.get("columns")
    latest = metadata.get("latest_timestamp")
    data_period = "Not available"
    col_count = len(columns) if isinstance(columns, list) else None

    if cleaned_xlsx.exists():
        try:
            df = pd.read_excel(cleaned_xlsx, usecols=lambda c: c in ["date", "time"])
            rows = rows or len(df)
            if col_count is None:
                col_count = len(pd.read_excel(cleaned_xlsx, nrows=0).columns)
            if not df.empty and "date" in df.columns:
                dates = pd.to_datetime(df["date"], errors="coerce")
                if dates.notna().any():
                    data_period = f"{dates.min().date()} to {dates.max().date()}"
                    latest = latest or str(dates.max())
        except Exception as exc:
            warn(f"Could not infer cleaned dataset summary: {exc}")
    else:
        warn(f"Missing source file: {rel(cleaned_xlsx)}")

    summary = pd.DataFrame(
        [
            ["Data period", data_period],
            ["Number of rows", rows if rows is not None else "Not available"],
            ["Number of columns", col_count if col_count is not None else "Not available"],
            ["Latest timestamp", latest if latest else "Not available"],
            ["Cleaned Excel path", rel(cleaned_xlsx)],
            ["Cleaned Parquet path", rel(cleaned_parquet)],
            ["Planned outage path", rel(outage_path)],
        ],
        columns=["Item", "Value"],
    )
    save_table(summary, "appendix_A_cleaned_dataset_summary", csv=False)
    return rules, summary


def appendix_b() -> pd.DataFrame:
    groups = pd.DataFrame(
        [
            ["Temporal features", "hour_sin, hour_cos, day_sin, day_cos, month_sin, month_cos, is_weekend", "Calendar and time-position indicators.", "RBFNN, Random Forest, XGBoost"],
            ["Target-hour cyclical features", "target_hour_sin, target_hour_cos, target_day_sin, target_day_cos, target_is_weekend", "Cyclical indicators for the forecasted hour.", "RBFNN"],
            ["Current plant generation", "total_gen_<plant>_current", "Current plant total generation used as forecast anchor.", "RBFNN, Random Forest, XGBoost"],
            ["Lag features", "lags at 1, 2, 3, 6, 12, 24, 48, 72, and 168 hours", "Historical plant generation at fixed lookback periods.", "RBFNN, Random Forest, XGBoost"],
            ["Rolling statistics", "rolling mean, standard deviation, minimum, and maximum over 3, 6, 12, 24, 48, and 168 hours", "Recent and weekly generation distribution behavior.", "RBFNN"],
            ["Generation difference features", "diff1, diff3, diff24", "Recent generation changes and ramp behavior.", "RBFNN"],
            ["Same-hour historical target references", "target_lag24, target_lag168", "Same-hour yesterday and same-hour previous-week references.", "RBFNN"],
            ["Unit-level generation features", "gen_<plant>_unit<n>", "Cleaned unit-level generation variables.", "RBFNN, Random Forest, XGBoost"],
            ["Unit share features", "unit generation share patterns", "Relative unit contribution to plant generation.", "RBFNN, Random Forest, XGBoost"],
            ["Unit outage/status features", "out_<plant>_unit<n>", "Binary unit availability or outage status.", "RBFNN, Random Forest, XGBoost"],
            ["Running units and plant availability", "<plant>_units_running, <plant>_plant_available", "Counts and flags describing plant-level availability.", "RBFNN, Random Forest, XGBoost"],
            ["Gate features", "tot_<plant>_gate and lagged gate features", "Gate opening and discharge control indicators.", "RBFNN"],
            ["Reservoir elevation features", "elev_<plant> and lagged elevation features", "Reservoir or forebay elevation status.", "RBFNN"],
            ["Spillway features", "tot_<plant> and lagged spillway/outflow features", "Plant water release and spillway-related variables.", "RBFNN"],
            ["Rainfall features", "rainfall and lagged rainfall features", "Hydrologic inflow proxy from rainfall measurements.", "RBFNN"],
            ["Lake Lanao outflow features", "lake_lanao_outflow and lagged outflow features", "Upstream lake outflow indicators.", "RBFNN"],
            ["Upstream cascade generation features", "upstream total generation features by cascade order", "Generation context from upstream plants.", "RBFNN"],
        ],
        columns=["Feature Group", "Feature Name or Pattern", "Description", "Used By"],
    )
    save_table(groups, "appendix_B_complete_feature_list")

    sheets = {"Feature Groups": groups}
    for key, name in zip(PLANT_KEYS, PLANTS):
        meta = read_json_safe(THESIS_DIR / f"models/rbfnn/meta_{key}.json")
        features = meta.get("X_cols") or meta.get("features") or []
        if features:
            sheets[f"{name} Features"] = pd.DataFrame({"Feature Name": features})
    if len(sheets) > 1:
        save_workbook(TABLES_DIR / "appendix_B_complete_feature_list_by_plant.xlsx", sheets)
    return groups


def appendix_c() -> pd.DataFrame:
    defaults = {
        "Agus 1": [158, 80, 0.001, 80, 32, 1.15, "Yes", "Bin calibration", "Adam", "Mean Squared Error"],
        "Agus 2": [169, 120, 0.0005, 80, 32, 1.15, "Yes", "None specified", "Adam", "Mean Squared Error"],
        "Agus 4": [169, 180, 0.0005, 80, 32, 1.15, "Yes", "None specified", "Adam", "Mean Squared Error"],
        "Agus 5": [168, 180, 0.001, 80, 32, 1.15, "Yes", "Bin calibration and hourly residual correction", "Adam", "Mean Squared Error"],
        "Agus 6": [171, 80, 0.0005, 80, 32, 1.15, "Yes", "None specified", "Adam", "Mean Squared Error"],
        "Agus 7": [168, 120, 0.001, 80, 32, 1.15, "Yes", "Hourly residual correction", "Adam", "Mean Squared Error"],
    }
    rbfnn_rows = []
    for key, plant in zip(PLANT_KEYS, PLANTS):
        feature_count, centers, lr, epochs, batch, shrink, bias, extra, opt, loss = defaults[plant]
        meta = read_json_safe(THESIS_DIR / f"models/rbfnn/meta_{key}.json")
        feature_count = len(meta.get("X_cols", [])) or meta.get("feature_count", feature_count)
        centers = meta.get("selected_centers", meta.get("centers", meta.get("rbf_centers", centers)))
        lr = meta.get("selected_learning_rate", meta.get("learning_rate", lr))
        shrink = meta.get("best_shrinkage", meta.get("shrinkage", shrink))
        bias_value = meta.get("bias_correction_mw")
        if bias_value is not None:
            bias = f"Yes ({bias_value:.6f} MW)"
        rbfnn_rows.append([plant, feature_count, centers, lr, epochs, batch, shrink, bias, extra, opt, loss])

    rbfnn_df = pd.DataFrame(
        rbfnn_rows,
        columns=[
            "Plant",
            "Feature Count",
            "RBF Centers",
            "Learning Rate",
            "Epochs",
            "Batch Size",
            "Shrinkage",
            "Bias Correction",
            "Extra Calibration",
            "Optimizer",
            "Loss Function",
        ],
    )

    benchmark_rows = [
        ["Random Forest", "All plants", "n_estimators", 400, "Number of trees in the ensemble.", "scripts/cell3_benchmark.py"],
        ["Random Forest", "All plants", "min_samples_leaf", 2, "Minimum samples required at a leaf node.", "scripts/cell3_benchmark.py"],
        ["Random Forest", "All plants", "random_state", 42, "Reproducibility seed.", "scripts/cell3_benchmark.py"],
        ["Random Forest", "All plants", "n_jobs", 1, "Single-worker execution for reproducible local runs.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "n_estimators", 500, "Number of boosted trees.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "learning_rate", 0.03, "Step size for boosting updates.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "max_depth", 4, "Maximum tree depth.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "subsample", 0.9, "Row sampling fraction per tree.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "colsample_bytree", 0.9, "Column sampling fraction per tree.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "objective", "reg:squarederror", "Squared-error regression objective.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "random_state", 42, "Reproducibility seed.", "scripts/cell3_benchmark.py"],
        ["XGBoost", "All plants", "n_jobs", 1, "Single-worker execution for reproducible local runs.", "scripts/cell3_benchmark.py"],
    ]
    benchmark_df = pd.DataFrame(
        benchmark_rows,
        columns=["Model", "Scope", "Parameter", "Value", "Purpose", "Source File"],
    )

    combined_rows = []
    for _, row in rbfnn_df.iterrows():
        combined_rows.extend(
            [
                ["RBFNN", row["Plant"], "Feature Count", row["Feature Count"], "Number of input features used by the plant model.", f"models/rbfnn/meta_{row['Plant'].lower().replace(' ', '')}.json"],
                ["RBFNN", row["Plant"], "RBF Centers", row["RBF Centers"], "Selected number of radial basis centers.", f"models/rbfnn/meta_{row['Plant'].lower().replace(' ', '')}.json"],
                ["RBFNN", row["Plant"], "Learning Rate", row["Learning Rate"], "Adam optimizer learning rate selected during validation tuning.", f"models/rbfnn/meta_{row['Plant'].lower().replace(' ', '')}.json"],
                ["RBFNN", row["Plant"], "Epochs", row["Epochs"], "Maximum training epochs used by the RBFNN training script.", "scripts/cell2_rbfnn.py"],
                ["RBFNN", row["Plant"], "Batch Size", row["Batch Size"], "Mini-batch size used during model fitting.", "scripts/cell2_rbfnn.py"],
                ["RBFNN", row["Plant"], "Shrinkage", row["Shrinkage"], "Validation-selected shrinkage applied to residual forecasts.", f"models/rbfnn/meta_{row['Plant'].lower().replace(' ', '')}.json"],
                ["RBFNN", row["Plant"], "Bias Correction", row["Bias Correction"], "Bias correction applied during level prediction.", f"models/rbfnn/meta_{row['Plant'].lower().replace(' ', '')}.json"],
                ["RBFNN", row["Plant"], "Extra Calibration", row["Extra Calibration"], "Additional calibration noted for the selected plant configuration.", "thesis appendix configuration"],
                ["RBFNN", row["Plant"], "Optimizer", row["Optimizer"], "Optimizer used to train the neural network.", "scripts/cell2_rbfnn.py"],
                ["RBFNN", row["Plant"], "Loss Function", row["Loss Function"], "Training loss used for model fitting.", "scripts/cell2_rbfnn.py"],
            ]
        )
    combined_df = pd.concat(
        [
            pd.DataFrame(combined_rows, columns=["Model", "Scope", "Parameter", "Value", "Purpose", "Source File"]),
            benchmark_df,
        ],
        ignore_index=True,
    )

    save_table(rbfnn_df, "appendix_C_rbfnn_hyperparameters")
    save_table(combined_df, "appendix_C_model_hyperparameters_configuration")
    save_workbook(
        TABLES_DIR / "appendix_C_hyperparameters_configuration_all_models.xlsx",
        {
            "All Model Hyperparameters": combined_df,
            "RBFNN By Plant": rbfnn_df,
            "Random Forest XGBoost": benchmark_df,
        },
    )
    return combined_df


def appendix_d() -> pd.DataFrame:
    rows = [
        ["Random Forest", "n_estimators", 400, "Number of trees in the ensemble."],
        ["Random Forest", "min_samples_leaf", 2, "Minimum samples required at a leaf node."],
        ["Random Forest", "random_state", 42, "Reproducibility seed."],
        ["Random Forest", "n_jobs", 1, "Single-worker execution for reproducible local runs."],
        ["XGBoost", "n_estimators", 500, "Number of boosted trees."],
        ["XGBoost", "learning_rate", 0.03, "Step size for boosting updates."],
        ["XGBoost", "max_depth", 4, "Maximum tree depth."],
        ["XGBoost", "subsample", 0.9, "Row sampling fraction per tree."],
        ["XGBoost", "colsample_bytree", 0.9, "Column sampling fraction per tree."],
        ["XGBoost", "objective", "reg:squarederror", "Squared-error regression objective."],
        ["XGBoost", "random_state", 42, "Reproducibility seed."],
        ["XGBoost", "n_jobs", 1, "Single-worker execution for reproducible local runs."],
    ]
    df = pd.DataFrame(rows, columns=["Model", "Parameter", "Value", "Purpose"])
    save_table(df, "appendix_D_benchmark_configuration")
    return df


def appendix_e() -> tuple[pd.DataFrame, pd.DataFrame]:
    path = THESIS_DIR / "outputs/outages_planning/Planned_Outages_Input.xlsx"
    df = read_excel_safe(path).head(24)
    if not df.empty:
        save_table(df, "appendix_E_planned_outage_template")
        make_table_png(df, SCREENSHOTS_DIR / "appendix_E_planned_outage_template.png")
    legend = pd.DataFrame(
        [[1, "Unit available/on", "The unit is treated as available for forecasting."], [0, "Unit unavailable/outage", "The unit is excluded or treated as unavailable in the forecast horizon."]],
        columns=["Value", "Meaning", "Forecasting Interpretation"],
    )
    save_table(legend, "appendix_E_outage_status_legend", csv=False)
    return df, legend


def appendix_f() -> tuple[pd.DataFrame, pd.DataFrame]:
    path = THESIS_DIR / "outputs/rbfnn_forecast/Day_Ahead_24H_RBFNN_Forecast.xlsx"
    df = read_excel_safe(path).head(24)
    if not df.empty:
        save_table(df, "appendix_F_day_ahead_forecast_sample")
        make_table_png(df, SCREENSHOTS_DIR / "appendix_F_day_ahead_forecast_output.png")

    rows = []
    for plant in PLANTS:
        compact = plant.replace(" ", "")
        col = f"Total_Gen_{compact}_MW"
        if col in df.columns:
            values = pd.to_numeric(df[col], errors="coerce")
            rows.append([plant, values.sum(), values.mean(), values.max(), values.min()])
    cascade_col = "Total_Cascade_Generation_MW"
    if cascade_col in df.columns:
        values = pd.to_numeric(df[cascade_col], errors="coerce")
        rows.append(["Total Cascade", values.sum(), values.mean(), values.max(), values.min()])

    summary = pd.DataFrame(
        rows,
        columns=[
            "Plant",
            "Total Daily Forecasted Generation",
            "Average Hourly Forecast",
            "Maximum Hourly Forecast",
            "Minimum Hourly Forecast",
        ],
    )
    save_table(summary, "appendix_F_daily_total_generation_summary", csv=False)
    return df, summary


def appendix_g() -> pd.DataFrame:
    readme = SCREENSHOTS_DIR / "gui_reference/README_GUI_SCREENSHOTS.txt"
    readme.write_text(
        "Place GUI dashboard screenshots in this folder for Appendix G.\n\n"
        "Recommended screenshots:\n"
        "- Overview page\n"
        "- Data Management page\n"
        "- Planned Outage Planning page\n"
        "- Forecasting page\n"
        "- System Operation page\n\n"
        "Suggested filenames are listed in tables/appendix_G_gui_captions.xlsx.\n",
        encoding="utf-8",
    )
    GENERATED.append(readme)
    df = pd.DataFrame(
        [
            ["Figure G.1", "Dashboard overview page", "figure_G_1_dashboard_overview.png", "Dashboard overview page of the AI-based hydroelectric generation forecasting system."],
            ["Figure G.2", "Data management page", "figure_G_2_data_management.png", "Data management page for source files, cleaned data, and generated outputs."],
            ["Figure G.3", "Planned outage planning page", "figure_G_3_planned_outage_planning.png", "Planned outage planning page for entering unit availability over the forecast horizon."],
            ["Figure G.4", "Forecasting results page", "figure_G_4_forecasting_results.png", "Forecasting results page showing day-ahead generation forecasts."],
            ["Figure G.5", "System operation guide page", "figure_G_5_system_operation_guide.png", "System operation guide page containing execution guidance for users."],
        ],
        columns=["Figure Number", "Dashboard Page", "Suggested Screenshot Filename", "Caption"],
    )
    save_table(df, "appendix_G_gui_captions", csv=False)
    return df


def appendix_h() -> pd.DataFrame:
    df = pd.DataFrame(
        [
            ["Full process", "Clean and preprocess raw hydroelectric data.", "python scripts/cell1_clean_data.py", "Run when source data changes or cleaned outputs must be rebuilt."],
            ["Retraining mode", "Train and evaluate RBFNN forecasting models.", "python scripts/cell2_rbfnn.py --train", "Run when retraining RBFNN models is required."],
            ["Retraining mode", "Train and evaluate Random Forest and XGBoost benchmarks.", "python scripts/cell3_benchmark.py --train", "Run when benchmark models must be regenerated."],
            ["Full process", "Generate thesis metrics, figures, and supporting outputs.", "python scripts/cell4_generate_thesis_and_metrics.py", "Run after model outputs are updated."],
            ["Forecast-only mode", "Generate day-ahead RBFNN forecast using saved models.", "python scripts/cell2_rbfnn.py", "Run for routine forecast generation without retraining."],
            ["Forecast-only mode", "Generate benchmark forecasts or outputs using saved benchmark models.", "python scripts/cell3_benchmark.py", "Run for benchmark forecast-only operation."],
            ["GUI mode", "Launch the Streamlit dashboard.", "streamlit run streamlit_app.py", "Run when using the graphical dashboard."],
        ],
        columns=["Step", "Purpose", "Command", "When to Use"],
    )
    save_table(df, "appendix_H_execution_commands", csv=False)
    user_guide = LATEX_DIR / "appendix_H_user_guide.tex"
    user_guide.write_text(
        r"""\section{System User Guide and Execution Commands}
The forecasting workflow can be executed either as a full regeneration process, a retraining process, a forecast-only process, or through the graphical dashboard.

\begin{verbatim}
python scripts/cell1_clean_data.py
python scripts/cell2_rbfnn.py --train
python scripts/cell3_benchmark.py --train
python scripts/cell4_generate_thesis_and_metrics.py
python scripts/cell2_rbfnn.py
python scripts/cell3_benchmark.py
streamlit run streamlit_app.py
\end{verbatim}
""",
        encoding="utf-8",
    )
    GENERATED.append(user_guide)
    return df


def appendix_i() -> dict[str, pd.DataFrame]:
    sources = {
        "RBFNN Metrics": THESIS_DIR / "metadata/overall_metrics/rbfnn_validation_testing_metrics.xlsx",
        "Random Forest Metrics": THESIS_DIR / "metadata/overall_metrics/random_forest_validation_testing_metrics.xlsx",
        "XGBoost Metrics": THESIS_DIR / "metadata/overall_metrics/xgboost_validation_testing_metrics.xlsx",
    }
    sheets = {}
    for name, path in sources.items():
        df = read_excel_safe(path)
        sheets[name] = df
    comparison = pd.concat([df for df in sheets.values() if not df.empty], ignore_index=True)
    sheets["Overall Comparison"] = comparison
    save_workbook(TABLES_DIR / "appendix_I_testing_metrics_summary.xlsx", sheets)

    target = FIGURES_DIR / "appendix_I_testing_plots"
    for source_dir in [THESIS_DIR / "thesis_figures", THESIS_DIR / "outputs/testing_plots"]:
        if not source_dir.exists():
            warn(f"Missing plot source directory: {rel(source_dir)}")
            continue
        for plot in source_dir.rglob("*.png"):
            dest = target / plot.relative_to(source_dir)
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(plot, dest)
            GENERATED.append(dest)
    return sheets


def picture_caption(path: Path) -> tuple[str, str, str]:
    relative = path.relative_to(APPENDIX_DIR)
    parts = relative.parts
    filename = path.stem.replace("_", " ")
    source = "Generated inside appendix_materials"
    appendix = "Appendix figure"

    if "appendix_I_testing_plots" in parts:
        idx = parts.index("appendix_I_testing_plots")
        category = parts[idx + 1] if len(parts) > idx + 1 else ""
        appendix = "Appendix I"
        source_rel = Path(*parts[idx + 1 :])
        thesis_source = THESIS_DIR / "thesis_figures" / source_rel
        testing_source = THESIS_DIR / "outputs/testing_plots" / source_rel
        if thesis_source.exists():
            source = rel(thesis_source)
        elif testing_source.exists():
            source = rel(testing_source)
        else:
            source = "Copied from available thesis/testing plot source"

        category_captions = {
            "benchmark_actual_testing_plots": "Actual versus benchmark testing forecast plot",
            "benchmark_comparison": "Benchmark model testing metric comparison plot",
            "cleaned_profiles": "Cleaned generation profile plot",
            "day_ahead_forecast": "Day-ahead forecast plot",
            "error_analysis": "RBFNN testing error analysis plot",
            "full_testing_split": "Full testing split forecast comparison plot",
            "testing_actual_vs_forecast": "RBFNN testing actual versus forecast plot",
            "testing_metrics": "RBFNN testing metric plot",
            "validation_metrics": "RBFNN validation metric plot",
        }
        caption = category_captions.get(category, "Appendix testing or forecasting plot")
        plant = next((p for p in ["agus1", "agus2", "agus4", "agus5", "agus6", "agus7"] if p in path.name.lower()), "")
        if plant:
            caption = f"{caption} for {plant.replace('agus', 'Agus ')}"
        elif "all_plants" in path.name.lower():
            caption = f"{caption} for all Agus plants"
        return appendix, caption, source

    if path.name == "appendix_E_planned_outage_template.png":
        return "Appendix E", "Screenshot-style table image of the 24-hour planned outage input template.", "outputs/outages_planning/Planned_Outages_Input.xlsx"
    if path.name == "appendix_F_day_ahead_forecast_output.png":
        return "Appendix F", "Screenshot-style table image of the 24-hour day-ahead RBFNN forecast output.", "outputs/rbfnn_forecast/Day_Ahead_24H_RBFNN_Forecast.xlsx"

    return appendix, f"Appendix image file: {filename}.", source


def image_dimensions(path: Path) -> tuple[str, str]:
    try:
        from PIL import Image

        with Image.open(path) as image:
            width, height = image.size
        return str(width), str(height)
    except Exception:
        return "Not available", "Not available"


def make_picture_tables() -> pd.DataFrame:
    rows = []
    image_paths = sorted(list(FIGURES_DIR.rglob("*.png")) + list(SCREENSHOTS_DIR.rglob("*.png")))
    for index, image_path in enumerate(image_paths, start=1):
        appendix, caption, source = picture_caption(image_path)
        width, height = image_dimensions(image_path)
        row = {
            "Picture Number": f"P{index:03d}",
            "Appendix Section": appendix,
            "Picture Filename": image_path.name,
            "Appendix Picture Path": rel(image_path),
            "Source File or Workbook": source,
            "Figure Type": "PNG image",
            "Width Pixels": width,
            "Height Pixels": height,
            "File Size KB": round(image_path.stat().st_size / 1024, 2),
            "Factual Caption": caption,
            "Use in Thesis": "Use as an appendix figure or supporting visual reference.",
        }
        rows.append(row)

        individual = pd.DataFrame(
            [
                ["Picture Number", row["Picture Number"]],
                ["Appendix Section", row["Appendix Section"]],
                ["Picture Filename", row["Picture Filename"]],
                ["Appendix Picture Path", row["Appendix Picture Path"]],
                ["Source File or Workbook", row["Source File or Workbook"]],
                ["Figure Type", row["Figure Type"]],
                ["Width Pixels", row["Width Pixels"]],
                ["Height Pixels", row["Height Pixels"]],
                ["File Size KB", row["File Size KB"]],
                ["Factual Caption", row["Factual Caption"]],
                ["Use in Thesis", row["Use in Thesis"]],
            ],
            columns=["Field", "Value"],
        )
        safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in image_path.stem)[:90]
        out = TABLES_DIR / "picture_tables" / f"{row['Picture Number']}_{safe_name}.xlsx"
        individual.to_excel(out, index=False)
        format_workbook(out)
        GENERATED.append(out)

    catalog = pd.DataFrame(rows)
    save_table(catalog, "appendix_picture_catalog")
    save_workbook(
        TABLES_DIR / "appendix_picture_tables_by_section.xlsx",
        {
            "All Pictures": catalog,
            "Appendix E": catalog[catalog["Appendix Section"] == "Appendix E"],
            "Appendix F": catalog[catalog["Appendix Section"] == "Appendix F"],
            "Appendix I": catalog[catalog["Appendix Section"] == "Appendix I"],
        },
    )
    return catalog


def latex_appendices() -> None:
    path = LATEX_DIR / "appendices.tex"
    path.write_text(
        r"""\appendix

\chapter{Data Cleaning and Preprocessing Artifacts}
\section{Data Cleaning Rules}
\begin{table}[H]
\centering
\caption{Summary of data cleaning and preprocessing rules}
\label{tab:appendix_a_cleaning_rules}
\resizebox{\textwidth}{!}{%
\begin{tabular}{llll}
\hline
Cleaning Step & Rule Applied & Purpose & Source/Output File \\
\hline
\multicolumn{4}{l}{See \texttt{tables/appendix\_A\_data\_cleaning\_rules.xlsx}.} \\
\hline
\end{tabular}}
\end{table}

\chapter{Complete Forecasting Feature List}
\section{Feature Groups}
\begin{table}[H]
\centering
\caption{Complete forecasting feature groups}
\label{tab:appendix_b_feature_groups}
\resizebox{\textwidth}{!}{%
\begin{tabular}{llll}
\hline
Feature Group & Feature Name or Pattern & Description & Used By \\
\hline
\multicolumn{4}{l}{See \texttt{tables/appendix\_B\_complete\_feature\_list.xlsx}.} \\
\hline
\end{tabular}}
\end{table}

\chapter{Model Hyperparameter Configuration}
\begin{table}[H]
\centering
\caption{Hyperparameter configuration for RBFNN, Random Forest, and XGBoost models}
\label{tab:appendix_c_model_hyperparameters}
\resizebox{\textwidth}{!}{%
\begin{tabular}{llllll}
\hline
Model & Scope & Parameter & Value & Purpose & Source File \\
\hline
\multicolumn{6}{l}{See \texttt{tables/appendix\_C\_hyperparameters\_configuration\_all\_models.xlsx}.} \\
\hline
\end{tabular}}
\end{table}

\chapter{Benchmark Model Configuration}
\begin{table}[H]
\centering
\caption{Benchmark model hyperparameter configuration}
\label{tab:appendix_d_benchmark_configuration}
\resizebox{\textwidth}{!}{%
\begin{tabular}{llll}
\hline
Model & Parameter & Value & Purpose \\
\hline
\multicolumn{4}{l}{See \texttt{tables/appendix\_D\_benchmark\_configuration.xlsx}.} \\
\hline
\end{tabular}}
\end{table}

\chapter{Planned Outage Input Template}
\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/appendix_E_planned_outage_template.png}
\caption{Planned outage input template for the 24-hour forecast horizon}
\label{fig:appendix_e_planned_outage_template}
\end{figure}

\chapter{Sample Day-Ahead Forecast Output}
\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/appendix_F_day_ahead_forecast_output.png}
\caption{Sample 24-hour day-ahead RBFNN forecast output}
\label{fig:appendix_f_day_ahead_forecast_output}
\end{figure}

\chapter{Graphical User Interface Dashboard Reference}
\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/gui_reference/figure_G_1_dashboard_overview.png}
\caption{Dashboard overview page}
\label{fig:appendix_g_dashboard_overview}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/gui_reference/figure_G_2_data_management.png}
\caption{Data management page}
\label{fig:appendix_g_data_management}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/gui_reference/figure_G_3_planned_outage_planning.png}
\caption{Planned outage planning page}
\label{fig:appendix_g_planned_outage}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/gui_reference/figure_G_4_forecasting_results.png}
\caption{Forecasting results page}
\label{fig:appendix_g_forecasting_results}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/screenshots/gui_reference/figure_G_5_system_operation_guide.png}
\caption{System operation guide page}
\label{fig:appendix_g_system_operation}
\end{figure}

\chapter{System User Guide and Execution Commands}
\section{Execution Commands}
\begin{verbatim}
python scripts/cell1_clean_data.py
python scripts/cell2_rbfnn.py --train
python scripts/cell3_benchmark.py --train
python scripts/cell4_generate_thesis_and_metrics.py
python scripts/cell2_rbfnn.py
python scripts/cell3_benchmark.py
streamlit run streamlit_app.py
\end{verbatim}

\chapter{Additional Testing Results and Forecasting Plots}
\section{Testing Metrics and Forecasting Figures}
\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/figures/appendix_I_testing_plots/benchmark_comparison/testing_mape_model_comparison.png}
\caption{Testing MAPE comparison across forecasting models}
\label{fig:appendix_i_testing_mape_comparison}
\end{figure}

\begin{figure}[H]
\centering
\includegraphics[width=\textwidth]{appendix_materials/figures/appendix_I_testing_plots/benchmark_comparison/testing_rmse_model_comparison.png}
\caption{Testing RMSE comparison across forecasting models}
\label{fig:appendix_i_testing_rmse_comparison}
\end{figure}
""",
        encoding="utf-8",
    )
    GENERATED.append(path)


def appendix_summary(
    rules: pd.DataFrame,
    feature_groups: pd.DataFrame,
    model_hyper: pd.DataFrame,
    benchmark: pd.DataFrame,
    outage_legend: pd.DataFrame,
    commands: pd.DataFrame,
    testing_sheets: dict[str, pd.DataFrame],
    picture_catalog: pd.DataFrame,
) -> None:
    appendix_list = pd.DataFrame(
        [
            ["Appendix A", "Data Cleaning and Preprocessing Artifacts", "Cleaning rules and cleaned dataset summary."],
            ["Appendix B", "Complete Forecasting Feature List", "Feature groups and per-plant RBFNN feature sheets where available."],
            ["Appendix C", "Model Hyperparameter Configuration", "RBFNN, Random Forest, and XGBoost hyperparameter settings."],
            ["Appendix D", "Benchmark Model Configuration", "Random Forest and XGBoost parameter settings."],
            ["Appendix E", "Planned Outage Input Template", "24-hour planned outage input and unit status legend."],
            ["Appendix F", "Sample Day-Ahead Forecast Output", "24-hour forecast sample and daily summary."],
            ["Appendix G", "GUI Dashboard Reference", "Dashboard screenshot captions and placement instructions."],
            ["Appendix H", "System User Guide and Execution Commands", "Full process, retraining, forecast-only, and GUI commands."],
            ["Appendix I", "Additional Testing Results and Forecasting Plots", "Testing metrics and copied plot references."],
        ],
        columns=["Appendix", "Title", "Contents"],
    )
    testing_summary = testing_sheets.get("Overall Comparison", pd.DataFrame())
    save_workbook(
        APPENDIX_DIR / "appendix_summary.xlsx",
        {
            "Appendix List": appendix_list,
            "Data Cleaning Rules": rules,
            "Feature Groups": feature_groups,
            "Model Hyperparameters": model_hyper,
            "Benchmark Configuration": benchmark,
            "Outage Legend": outage_legend,
            "Execution Commands": commands,
            "Testing Metrics Summary": testing_summary,
            "Picture Catalog": picture_catalog,
        },
    )


def main() -> None:
    ensure_dirs()

    main_sources = [
        THESIS_DIR / "outputs/cleaned_data/cleaned_hourly_data.xlsx",
        THESIS_DIR / "outputs/outages_planning/Planned_Outages_Input.xlsx",
        THESIS_DIR / "outputs/rbfnn_forecast/Day_Ahead_24H_RBFNN_Forecast.xlsx",
        THESIS_DIR / "metadata/overall_metrics/rbfnn_validation_testing_metrics.xlsx",
    ]
    if not any(path.exists() for path in main_sources):
        raise FileNotFoundError("No main appendix source files were found in Thesis Forecasting.")

    rules, _cleaned_summary = appendix_a()
    feature_groups = appendix_b()
    model_hyper = appendix_c()
    benchmark = appendix_d()
    _outage, outage_legend = appendix_e()
    appendix_f()
    appendix_g()
    commands = appendix_h()
    testing_sheets = appendix_i()
    picture_catalog = make_picture_tables()
    latex_appendices()
    appendix_summary(rules, feature_groups, model_hyper, benchmark, outage_legend, commands, testing_sheets, picture_catalog)

    print("\nGenerated appendix materials:")
    for path in sorted(set(GENERATED)):
        print(f"- {rel(path)}")
    if WARNINGS:
        print("\nWarnings:")
        for message in WARNINGS:
            print(f"- {message}")
    print(f"\nTotal generated files: {len(set(GENERATED))}")


if __name__ == "__main__":
    main()
