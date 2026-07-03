"""
Optimized Cell 3 benchmarks.

Random Forest and XGBoost are comparison models only. They use the same
chronological split and operational MAPE definition as optimized Cell 2.
"""

import json
import math
import re
import sys
import warnings
import zipfile
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from xgboost import XGBRegressor

warnings.filterwarnings("ignore", category=pd.errors.PerformanceWarning)
warnings.filterwarnings("ignore", message="`sklearn.utils.parallel.delayed` should be used*")

# ============================================================
# PATH AND BENCHMARK CONFIGURATION
# ============================================================

# Defines project paths for cleaned data, benchmark models, metrics, and forecast outputs.
PROJECT_DIR = Path(__file__).resolve().parents[2]
THESIS_DIR = Path(__file__).resolve().parents[1]
OPT_DIR = THESIS_DIR
LEGACY_OPT_DIR = PROJECT_DIR / "archive" / "optimized_version"
OUT_DIR = OPT_DIR / "outputs"
CLEANED_DATA_DIR = OUT_DIR / "cleaned_data"
OUTAGES_DIR = OUT_DIR / "outages_planning"
BENCHMARK_OUTPUT_DIRS = {
    "random_forest": OUT_DIR / "random_forest_forecast",
    "xgboost": OUT_DIR / "xgboost_forecast",
}
MODEL_DIRS = {
    "random_forest": OPT_DIR / "models" / "random_forest",
    "xgboost": OPT_DIR / "models" / "xgboost",
}
LEGACY_MODEL_DIRS = {
    "random_forest": LEGACY_OPT_DIR / "models" / "random_forest",
    "xgboost": LEGACY_OPT_DIR / "models" / "xgboost",
}
META_DIR = OPT_DIR / "metadata" / "overall_metrics"
DAY_AHEAD_BACKTEST_DIR = OPT_DIR / "metadata" / "day_ahead_backtest"

META_DIR.mkdir(parents=True, exist_ok=True)
for folder in [CLEANED_DATA_DIR, OUTAGES_DIR, DAY_AHEAD_BACKTEST_DIR, *BENCHMARK_OUTPUT_DIRS.values(), *MODEL_DIRS.values()]:
    folder.mkdir(parents=True, exist_ok=True)

# Stores plant capacities, output schemas, feature windows, and benchmark model keys.
PLANTS = ["agus1", "agus2", "agus4", "agus5", "agus6", "agus7"]
UPSTREAM_MAP = {"agus1": None, "agus2": "agus1", "agus4": "agus2", "agus5": "agus4", "agus6": "agus5", "agus7": "agus6"}
METRICS_COLUMNS = [
    "model",
    "plant",
    "feature_count",
    "val_operational_mape",
    "val_mae",
    "val_rmse",
    "val_r2",
    "test_operational_mape",
    "test_mae",
    "test_rmse",
    "test_r2",
]
TESTING_PREDICTION_COLUMNS = [
    "Date",
    "Hour",
    "datetime",
    "plant",
    "actual_generation",
    "predicted_generation",
    "model",
]
CAPACITY_MW = {"agus1": 80.0, "agus2": 180.0, "agus4": 158.1, "agus5": 55.0, "agus6": 219.0, "agus7": 54.0}
UNIT_CAPACITY = {
    "agus1": {"unit1": 40.0, "unit2": 40.0},
    "agus2": {"unit1": 60.0, "unit2": 60.0, "unit3": 60.0},
    "agus4": {"unit1": 52.7, "unit2": 52.7, "unit3": 52.7},
    "agus5": {"unit1": 27.5, "unit2": 27.5},
    "agus6": {"unit1": 34.5, "unit2": 34.5, "unit3": 50.0, "unit4": 50.0, "unit5": 50.0},
    "agus7": {"unit1": 27.0, "unit2": 27.0},
}
UNIT_FORECAST_COLUMNS = [
    f"gen_{plant}_{unit}"
    for plant in PLANTS
    for unit in UNIT_CAPACITY[plant]
]
TOTAL_FORECAST_COLUMNS = [f"total_gen_{plant}" for plant in PLANTS]
CASCADE_FORECAST_COLUMN = "total_cascade_generation"
LAGS = [1, 2, 3, 6, 12, 24, 48, 72, 168]
ROLL_WINDOWS = [3, 6, 12, 24, 48, 168]
MODEL_KEYS = {"Random Forest": "random_forest", "XGBoost": "xgboost"}
FEATURE_HISTORY_WINDOW = 240
BENCHMARK_FINALIST_COUNT = 3
SELECTION_DAY_AHEAD_MAX_DAYS = 12
R2_WEIGHT = 8.0
NEGATIVE_R2_PENALTY = 25.0
DAY_AHEAD_EVALUATION_TYPE = "rolling_24h_day_ahead_backtest"
OUTAGE_INPUT_TYPE = "last_known_outage_status_before_forecast_day"
OUTAGE_INPUT_NOTE = (
    "Validation/testing day-ahead backtests persist only unit outage/status columns "
    "from the latest historical row before each forecast day. Same-day actual unit "
    "outage/status values are not used because archived planned outage schedules "
    "were unavailable."
)
HYDROLOGIC_INPUT_TYPE = "lagged_or_persistence_context"
HYDROLOGIC_INPUT_NOTE = (
    "Forecast-day rainfall and Lake Lanao outflow were not taken from actual "
    "same-day records; lagged or most recent known values were used to avoid "
    "future leakage."
)
DAY_AHEAD_THESIS_NOTE = (
    "Each validation and testing day was forecast as a complete 24-hour horizon "
    "using only data available before the forecast day. Lag and rolling features "
    "were reconstructed from historical values and recursive predictions, while "
    "actual forecast-day generation was used only after prediction for metrics."
)


# ============================================================
# METRICS, FILE HELPERS, AND MODEL FACTORIES
# ============================================================

# Sets the minimum actual-generation level for meaningful MAPE calculation.
def operational_threshold(plant):
    return max(1.0, 0.01 * CAPACITY_MW[plant])


# Computes MAPE only on operational rows to avoid near-zero percentage distortion.
def operational_mape(y_true, y_pred, plant):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    mask = np.abs(y_true) >= operational_threshold(plant)
    if not mask.any():
        return np.nan
    return float(np.mean(np.abs((y_true[mask] - y_pred[mask]) / np.abs(y_true[mask]))) * 100.0)


# Collects benchmark validation/testing metrics for comparison with RBFNN.
def metrics(y_true, y_pred, plant):
    return {
        "operational_mape": operational_mape(y_true, y_pred, plant),
        "mae": mean_absolute_error(y_true, y_pred),
        "rmse": math.sqrt(mean_squared_error(y_true, y_pred)),
        "r2": r2_score(y_true, y_pred),
    }


def leakage_feature_audit(feature_columns):
    forbidden_patterns = [r"_tplus1$", r"_delta_tplus1$", r"actual", r"predicted", r"forecast", r"testing", r"validation"]
    forbidden = [
        col for col in feature_columns
        if any(re.search(pattern, col, flags=re.IGNORECASE) for pattern in forbidden_patterns)
    ]
    return "PASS" if not forbidden else "FAIL"


def level_from_delta(current, delta_pred, plant):
    return np.clip(
        np.asarray(current, dtype=float) + np.asarray(delta_pred, dtype=float),
        0.0,
        CAPACITY_MW[plant] * 1.05,
    )


SHRINKAGE_GRID = [0.05, 0.10, 0.20, 0.35, 0.50, 0.75, 1.00, 1.15]
BIN_CALIBRATION_QUANTILES = list(range(3, 21))
BIN_CALIBRATION_SCALES = [round(x, 2) for x in np.arange(0.50, 2.55, 0.05)]
PROFILE_BLEND_GRID = [0.0, 0.10, 0.20, 0.30, 0.40, 0.55, 0.70, 0.85]
HOURLY_CORRECTION_SCALE_GRID = [0.0, 0.25, 0.50, 0.75, 1.00, 1.25]
ANCHOR_BLEND_GRID = [0.0, 0.25, 0.50, 0.75, 1.00]
ANCHOR_CURRENT_WEIGHT_GRID = [0.0, 0.25, 0.50, 0.75, 1.00]


def calibrated_level_prediction(current, delta_pred, plant, calibration=None):
    calibration = calibration or {}
    shrinkage = float(calibration.get("shrinkage", 1.0))
    bias = float(calibration.get("bias_correction_mw", 0.0))
    return np.clip(
        np.asarray(current, dtype=float) + shrinkage * np.asarray(delta_pred, dtype=float) + bias,
        0.0,
        CAPACITY_MW[plant] * 1.05,
    )


def fit_shrinkage_bias(current, delta_pred, actual, plant):
    best = None
    for shrinkage in SHRINKAGE_GRID:
        raw_pred = np.clip(
            np.asarray(current, dtype=float) + shrinkage * np.asarray(delta_pred, dtype=float),
            0.0,
            CAPACITY_MW[plant] * 1.05,
        )
        bias = float(np.median(np.asarray(actual, dtype=float) - raw_pred))
        bias = float(np.clip(bias, -operational_threshold(plant), operational_threshold(plant)))
        pred = np.clip(raw_pred + bias, 0.0, CAPACITY_MW[plant] * 1.05)
        score = candidate_score(metrics(actual, pred, plant))
        if best is None or score < best["score"]:
            best = {
                "score": score,
                "calibration": {
                    "shrinkage": float(shrinkage),
                    "bias_correction_mw": bias,
                    "bin_calibration": None,
                    "hourly_residual_correction": None,
                    "profile_blend": None,
                },
                "pred": pred,
            }
    return best["calibration"], best["pred"]


def calibration_bins(values, q):
    _, bins = pd.qcut(pd.Series(values), q, duplicates="drop", retbins=True)
    return np.asarray(bins, dtype=float)


def bin_ids(values, bins):
    return np.digitize(np.asarray(values, dtype=float), bins[1:-1])


def apply_bin_calibration(pred, basis_values, plant, calibration):
    if not calibration:
        return np.asarray(pred, dtype=float)
    bins = np.asarray(calibration["bins"], dtype=float)
    corrections = {int(k): float(v) for k, v in calibration["corrections"].items()}
    scale = float(calibration["scale"])
    ids = bin_ids(basis_values, bins)
    out = np.asarray(pred, dtype=float).copy()
    for idx, bin_id in enumerate(ids):
        out[idx] += scale * corrections.get(int(bin_id), 0.0)
    return np.clip(out, 0.0, CAPACITY_MW[plant] * 1.05)


def tune_bin_calibration(plant, basis_values, actual, pred):
    best = {"score": candidate_score(metrics(actual, pred, plant)), "calibration": None, "pred": np.asarray(pred, dtype=float)}
    for q in BIN_CALIBRATION_QUANTILES:
        try:
            bins = calibration_bins(basis_values, q)
        except ValueError:
            continue
        ids = bin_ids(basis_values, bins)
        residual = np.asarray(actual, dtype=float) - np.asarray(pred, dtype=float)
        corrections = {
            int(bin_id): float(np.median(residual[ids == bin_id]))
            for bin_id in np.unique(ids)
            if np.any(ids == bin_id)
        }
        for scale in BIN_CALIBRATION_SCALES:
            calibration = {"bins": bins.tolist(), "corrections": corrections, "scale": float(scale), "basis": "current"}
            candidate = apply_bin_calibration(pred, basis_values, plant, calibration)
            score = candidate_score(metrics(actual, candidate, plant))
            if score < best["score"]:
                best = {"score": score, "calibration": calibration, "pred": candidate}
    return best["calibration"], best["pred"]


def target_hours_from_rows(df):
    return ((pd.to_numeric(df["time"]).astype(int) % 24) + 1).astype(int).values


def apply_hourly_correction(pred, hours, plant, correction):
    if not correction:
        return np.asarray(pred, dtype=float)
    scale = float(correction.get("scale", 0.0))
    by_hour = {int(k): float(v) for k, v in correction.get("hour_corrections", {}).items()}
    offsets = np.asarray([by_hour.get(int(hour), 0.0) for hour in hours], dtype=float)
    return np.clip(np.asarray(pred, dtype=float) + scale * offsets, 0.0, CAPACITY_MW[plant] * 1.05)


def tune_hourly_correction(plant, pred, actual, hours):
    residual = np.asarray(actual, dtype=float) - np.asarray(pred, dtype=float)
    hours = np.asarray(hours, dtype=int)
    hour_corrections = {
        int(hour): float(np.median(residual[hours == hour]))
        for hour in np.unique(hours)
        if np.any(hours == hour)
    }
    best = {"score": candidate_score(metrics(actual, pred, plant)), "correction": None, "pred": np.asarray(pred, dtype=float)}
    for scale in HOURLY_CORRECTION_SCALE_GRID:
        correction = {"scale": float(scale), "hour_corrections": hour_corrections}
        candidate = apply_hourly_correction(pred, hours, plant, correction)
        score = candidate_score(metrics(actual, candidate, plant))
        if score < best["score"]:
            best = {"score": score, "correction": correction, "pred": candidate}
    return best["correction"], best["pred"]


def same_hour_target_anchor(df, plant):
    target = f"total_gen_{plant}"
    col = f"{target}_target_lag24"
    if col in df.columns:
        return df[col].values.astype(float)
    return df[target].shift(23).values.astype(float)


def apply_profile_blend(pred, anchor, plant, config):
    if not config:
        return np.asarray(pred, dtype=float)
    weight = float(config.get("weight", 0.0))
    anchor = np.nan_to_num(np.asarray(anchor, dtype=float), nan=0.0)
    return np.clip((1.0 - weight) * np.asarray(pred, dtype=float) + weight * anchor, 0.0, CAPACITY_MW[plant] * 1.05)


def tune_profile_blend(plant, pred, actual, anchor):
    best = {"score": candidate_score(metrics(actual, pred, plant)), "config": None, "pred": np.asarray(pred, dtype=float)}
    for weight in PROFILE_BLEND_GRID:
        config = {"weight": float(weight)}
        candidate = apply_profile_blend(pred, anchor, plant, config)
        score = candidate_score(metrics(actual, candidate, plant))
        if score < best["score"]:
            best = {"score": score, "config": config, "pred": candidate}
    return best["config"], best["pred"]


def apply_benchmark_calibration(pred, basis_current, plant, calibration, hours=None, anchor=None):
    out = np.asarray(pred, dtype=float)
    if not calibration:
        return out
    out = apply_bin_calibration(out, basis_current, plant, calibration.get("bin_calibration"))
    if hours is not None:
        out = apply_hourly_correction(out, hours, plant, calibration.get("hourly_residual_correction"))
    if anchor is not None:
        out = apply_profile_blend(out, anchor, plant, calibration.get("profile_blend"))
    return np.clip(out, 0.0, CAPACITY_MW[plant] * 1.05)


def apply_level_anchor_blend(pred, current, lag_anchor, plant, config):
    if not config:
        return np.asarray(pred, dtype=float)
    weight = float(config.get("weight", 0.0))
    current_weight = float(config.get("current_weight", 0.0))
    current = np.nan_to_num(np.asarray(current, dtype=float), nan=0.0)
    lag_anchor = np.nan_to_num(np.asarray(lag_anchor, dtype=float), nan=0.0)
    anchor = current_weight * current + (1.0 - current_weight) * lag_anchor
    out = (1.0 - weight) * np.asarray(pred, dtype=float) + weight * anchor
    return np.clip(out, 0.0, CAPACITY_MW[plant] * 1.05)


def tune_level_anchor_blend(plant, pred, actual, current, lag_anchor):
    best = {"score": candidate_score(metrics(actual, pred, plant)), "config": None, "pred": np.asarray(pred, dtype=float)}
    for current_weight in ANCHOR_CURRENT_WEIGHT_GRID:
        for weight in ANCHOR_BLEND_GRID:
            config = {"weight": float(weight), "current_weight": float(current_weight)}
            candidate = apply_level_anchor_blend(pred, current, lag_anchor, plant, config)
            score = candidate_score(metrics(actual, candidate, plant))
            if score < best["score"]:
                best = {"score": score, "config": config, "pred": candidate}
    return best["config"], best["pred"]


def benchmark_level_prediction(current, delta_pred, plant, calibration=None, hours=None, anchor=None):
    calibration = calibration or {}
    out = calibrated_level_prediction(current, delta_pred, plant, calibration)
    out = apply_bin_calibration(out, current, plant, calibration.get("bin_calibration"))
    if hours is not None:
        out = apply_hourly_correction(out, hours, plant, calibration.get("hourly_residual_correction"))
    if anchor is not None:
        out = apply_profile_blend(out, anchor, plant, calibration.get("profile_blend"))
        out = apply_level_anchor_blend(out, current, anchor, plant, calibration.get("level_anchor_blend"))
    return np.clip(out, 0.0, CAPACITY_MW[plant] * 1.05)


def with_payload_calibration(model_info, calibration):
    model, model_feature_columns = unpack_model_payload(model_info)
    return model_payload(model, model_feature_columns, calibration)


# Formats Excel outputs for easier review and thesis reporting.
def format_excel(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            header = str(column_cells[0].value or "")
            max_len = max([len(header)] + [len(str(cell.value)) for cell in column_cells[1:80] if cell.value is not None])
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 34)
    wb.save(path)


def normalize_export_column_name(column):
    return str(column).strip().lower().replace(" ", "_").replace("-", "_")


def prepare_forecast_export(df):
    drop_columns = [
        col for col in df.columns
        if normalize_export_column_name(col) in {"forecast_hour", "datetime"}
    ]
    return df.drop(columns=drop_columns, errors="ignore")


def save_benchmark_fairness_audit(rows):
    path = META_DIR / "benchmark_fairness_audit.xlsx"
    columns = [
        "model",
        "plant",
        "feature count",
        "target type",
        "split strategy",
        "operational MAPE definition",
        "feature set alignment",
        "leakage check",
        "benchmark role",
    ]
    pd.DataFrame(rows)[columns].to_excel(path, index=False)
    format_excel(path)
    print("Saved:", path)


def update_model_selection_audit(section, rows):
    path = META_DIR / "model_selection_audit.json"
    audit = {}
    if path.exists():
        audit = json.loads(path.read_text())
    existing_rows = audit.get(section, [])
    if existing_rows and rows:
        merged = {
            (item.get("model"), item.get("plant")): item
            for item in existing_rows
        }
        for item in rows:
            merged[(item.get("model"), item.get("plant"))] = item
        audit[section] = list(merged.values())
    else:
        audit[section] = rows or existing_rows
    audit["controls"] = {
        "split_strategy": "chronological 70/15/15",
        "primary_model": "RBFNN",
        "benchmark_models": ["Random Forest", "XGBoost"],
        "selection_rule": "lowest R2-aware validation score; score penalizes MAPE, RMSE, low R2, and especially negative R2",
        "day_ahead_selection": f"models are selected with validation-only screening and R2-aware rolling 24-hour validation backtests sampled across up to {SELECTION_DAY_AHEAD_MAX_DAYS} validation days",
        "testing_used_for_tuning": False,
        "metric_source": "real model predictions only",
    }
    path.write_text(json.dumps(audit, indent=2, default=str))
    print("Saved:", path)


# Rebuilds the hourly datetime column from public date and time fields.
def rebuild_datetime(df):
    out = df.copy()
    out["date"] = pd.to_datetime(out["date"])
    out["time"] = pd.to_numeric(out["time"]).astype(int)
    if "datetime" not in out.columns:
        out["datetime"] = out["date"] + pd.to_timedelta(out["time"] - 1, unit="h")
    else:
        out["datetime"] = pd.to_datetime(out["datetime"])
    return out.sort_values("datetime").reset_index(drop=True)


# Adds legacy-compatible aliases expected by older saved benchmark models.
def add_runtime_compatibility_columns(df):
    # Keep Thesis Forecasting cleaned files public-facing while supporting models trained on old feature names.
    out = df.copy()
    for plant in PLANTS:
        new_gate = f"tot_{plant}_gate"
        old_gate = f"tot_{plant}"
        if new_gate in out.columns and old_gate not in out.columns:
            out[old_gate] = out[new_gate]
    if "lake_lanao_outflow" in out.columns and "lake_lanao_hourly_outflow_elev_agus7" not in out.columns:
        out["lake_lanao_hourly_outflow_elev_agus7"] = out["lake_lanao_outflow"]
    return out


# Resolves current benchmark model artifacts first, then legacy artifacts if needed.
def model_file(model_key, name):
    current = MODEL_DIRS[model_key] / name
    if current.exists():
        return current
    legacy = LEGACY_MODEL_DIRS[model_key] / name
    if legacy.exists():
        return legacy
    return current


# Aligns forecast features to the exact feature order saved with a model.
def align_features(X, model_features):
    X = X.copy()
    for col in model_features:
        if col not in X.columns:
            X[col] = 0
    return X[model_features]


# Saves each model together with the feature columns and validation-only calibration used during training.
def model_payload(model, model_feature_columns, calibration=None):
    return {
        "model": model,
        "feature_columns": list(model_feature_columns),
        "calibration": calibration or {},
    }


# Reads a saved model payload while supporting older plain-model files.
def unpack_model_payload(payload):
    if isinstance(payload, dict) and "model" in payload:
        model = payload["model"]
        if hasattr(model, "n_jobs"):
            model.n_jobs = 1
        return model, list(payload.get("feature_columns") or [])
    feature_names = getattr(payload, "feature_names_in_", None)
    model_features = list(feature_names) if feature_names is not None else []
    if hasattr(payload, "n_jobs"):
        payload.n_jobs = 1
    return payload, model_features


def model_calibration(payload):
    if isinstance(payload, dict):
        return payload.get("calibration") or {}
    return {}


# Persists one benchmark model and its feature list.
def save_model_payload(model_key, plant, model, model_feature_columns, calibration=None):
    joblib.dump(
        model_payload(model, model_feature_columns, calibration),
        MODEL_DIRS[model_key] / f"{model_key}_{plant}.pkl",
    )


def benchmark_search_configs(name, plant=None):
    if name == "Random Forest":
        configs = [
            {"n_estimators": 240, "max_depth": 8, "min_samples_leaf": 6, "max_features": "sqrt"},
            {"n_estimators": 300, "max_depth": 12, "min_samples_leaf": 4, "max_features": 0.55},
            {"n_estimators": 360, "max_depth": 18, "min_samples_leaf": 2, "max_features": 0.75},
            {"n_estimators": 420, "max_depth": None, "min_samples_leaf": 3, "max_features": 0.85},
        ]
        if plant == "agus2":
            configs.extend([
                {"n_estimators": 300, "max_depth": 6, "min_samples_leaf": 10, "max_features": "sqrt"},
                {"n_estimators": 360, "max_depth": 10, "min_samples_leaf": 12, "max_features": 0.40},
            ])
        return configs
    if name == "XGBoost":
        return [
            {"n_estimators": 260, "learning_rate": 0.05, "max_depth": 3, "subsample": 0.95, "colsample_bytree": 0.95, "reg_lambda": 1.0},
            {"n_estimators": 420, "learning_rate": 0.035, "max_depth": 3, "subsample": 0.90, "colsample_bytree": 0.90, "reg_lambda": 1.0},
            {"n_estimators": 560, "learning_rate": 0.025, "max_depth": 4, "subsample": 0.88, "colsample_bytree": 0.85, "reg_lambda": 1.5},
            {"n_estimators": 720, "learning_rate": 0.018, "max_depth": 4, "subsample": 0.85, "colsample_bytree": 0.80, "reg_lambda": 2.0},
        ]
    raise ValueError(f"Unsupported benchmark model: {name}")


# Creates the requested benchmark estimator with reproducible settings.
def make_benchmark_model(name, config=None):
    config = config or benchmark_search_configs(name)[1]
    if name == "Random Forest":
        return RandomForestRegressor(random_state=42, n_jobs=-1, **config)
    if name == "XGBoost":
        return XGBRegressor(
            objective="reg:squarederror",
            random_state=42,
            n_jobs=-1,
            **config,
        )
    raise ValueError(f"Unsupported benchmark model: {name}")


def candidate_score(metric_values):
    mape = metric_values["operational_mape"]
    mape_score = float(mape) if pd.notna(mape) else float("inf")
    r2 = float(metric_values["r2"]) if pd.notna(metric_values["r2"]) else -999.0
    r2_penalty = max(0.0, -r2) * NEGATIVE_R2_PENALTY
    combined_score = mape_score + metric_values["rmse"] + (1.0 - r2) * R2_WEIGHT + r2_penalty
    return (combined_score, -r2, mape_score, metric_values["rmse"])


# ============================================================
# BENCHMARK TRAINING AND PREDICTION HELPERS
# ============================================================

# Trains one plant-specific benchmark model and saves it for reuse.
def train_benchmark_model(name, plant, train, x_cols, y_col, val=None, raw_df=None, validation_forecast_days=None):
    X_train = train[x_cols].copy()
    model_feature_columns = X_train.columns.tolist()
    if val is None or raw_df is None or validation_forecast_days is None:
        model = make_benchmark_model(name)
        print(f"Training {name} benchmark for {plant}")
        model.fit(X_train, train[y_col])
        save_model_payload(MODEL_KEYS[name], plant, model, model_feature_columns)
        return model_payload(model, model_feature_columns), {}

    target = f"total_gen_{plant}"
    val_actual = val[f"{target}_tplus1"].values.astype(float)
    screened = []
    for config in benchmark_search_configs(name, plant):
        model = make_benchmark_model(name, config)
        print(f"Training {name} benchmark for {plant} with {config}")
        model.fit(X_train, train[y_col])
        payload = model_payload(model, model_feature_columns)
        val_delta = predict_aligned(payload, val[x_cols].copy(), plant)
        val_pred = level_from_delta(val[target], val_delta, plant)
        val_m = metrics(val_actual, val_pred, plant)
        screened.append({
            "screen_score": candidate_score(val_m),
            "screen_metrics": val_m,
            "config": config,
            "payload": payload,
        })

    screened = sorted(screened, key=lambda item: item["screen_score"])
    finalists = screened[:BENCHMARK_FINALIST_COUNT]
    best = None
    finalist_scores = []
    for candidate in finalists:
        rolling_score, rolling_metrics = rolling_benchmark_candidate_score(
            raw_df,
            validation_forecast_days,
            name,
            plant,
            candidate["payload"],
        )
        candidate["rolling_score"] = rolling_score
        candidate["rolling_validation_metrics"] = rolling_metrics
        finalist_scores.append({
            "config": candidate["config"],
            "one_step_validation_score": candidate["screen_score"],
            "rolling_validation_score": rolling_score,
            "rolling_validation_metrics": rolling_metrics,
        })
        if best is None or rolling_score < best["rolling_score"]:
            best = candidate

    model, _ = unpack_model_payload(best["payload"])
    val_delta = predict_aligned(best["payload"], val[x_cols].copy(), plant)
    calibration, val_pred = fit_shrinkage_bias(
        val[target].values.astype(float),
        val_delta,
        val_actual,
        plant,
    )
    calibration_payload = with_payload_calibration(best["payload"], calibration)
    calibrated_rolling_score, calibrated_rolling_metrics = rolling_benchmark_candidate_score(
        raw_df,
        validation_forecast_days,
        name,
        plant,
        calibration_payload,
    )
    if calibrated_rolling_score <= best["rolling_score"]:
        selected_calibration = calibration
        selected_rolling_score = calibrated_rolling_score
        selected_rolling_metrics = calibrated_rolling_metrics
    else:
        selected_calibration = {
            "shrinkage": 1.0,
            "bias_correction_mw": 0.0,
            "bin_calibration": None,
            "hourly_residual_correction": None,
            "profile_blend": None,
            "level_anchor_blend": None,
        }
        selected_rolling_score = best["rolling_score"]
        selected_rolling_metrics = best["rolling_validation_metrics"]
        val_pred = level_from_delta(val[target], val_delta, plant)

    bin_candidate, bin_val_pred = tune_bin_calibration(plant, val[target].values.astype(float), val_actual, val_pred)
    if bin_candidate:
        trial_calibration = dict(selected_calibration)
        trial_calibration["bin_calibration"] = bin_candidate
        trial_payload = with_payload_calibration(best["payload"], trial_calibration)
        trial_score, trial_metrics = rolling_benchmark_candidate_score(
            raw_df, validation_forecast_days, name, plant, trial_payload
        )
        if trial_score <= selected_rolling_score:
            selected_calibration = trial_calibration
            selected_rolling_score = trial_score
            selected_rolling_metrics = trial_metrics
            val_pred = bin_val_pred

    hourly_candidate, hourly_val_pred = tune_hourly_correction(plant, val_pred, val_actual, target_hours_from_rows(val))
    if hourly_candidate:
        trial_calibration = dict(selected_calibration)
        trial_calibration["hourly_residual_correction"] = hourly_candidate
        trial_payload = with_payload_calibration(best["payload"], trial_calibration)
        trial_score, trial_metrics = rolling_benchmark_candidate_score(
            raw_df, validation_forecast_days, name, plant, trial_payload
        )
        if trial_score <= selected_rolling_score:
            selected_calibration = trial_calibration
            selected_rolling_score = trial_score
            selected_rolling_metrics = trial_metrics
            val_pred = hourly_val_pred

    profile_candidate, profile_val_pred = tune_profile_blend(
        plant,
        val_pred,
        val_actual,
        same_hour_target_anchor(val, plant),
    )
    if profile_candidate:
        trial_calibration = dict(selected_calibration)
        trial_calibration["profile_blend"] = profile_candidate
        trial_payload = with_payload_calibration(best["payload"], trial_calibration)
        trial_score, trial_metrics = rolling_benchmark_candidate_score(
            raw_df, validation_forecast_days, name, plant, trial_payload
        )
        if trial_score <= selected_rolling_score:
            selected_calibration = trial_calibration
            selected_rolling_score = trial_score
            selected_rolling_metrics = trial_metrics

    anchor_candidate, anchor_val_pred = tune_level_anchor_blend(
        plant,
        val_pred,
        val_actual,
        val[target].values.astype(float),
        same_hour_target_anchor(val, plant),
    )
    if anchor_candidate:
        trial_calibration = dict(selected_calibration)
        trial_calibration["level_anchor_blend"] = anchor_candidate
        trial_payload = with_payload_calibration(best["payload"], trial_calibration)
        trial_score, trial_metrics = rolling_benchmark_candidate_score(
            raw_df, validation_forecast_days, name, plant, trial_payload
        )
        if trial_score <= selected_rolling_score:
            selected_calibration = trial_calibration
            selected_rolling_score = trial_score
            selected_rolling_metrics = trial_metrics
            val_pred = anchor_val_pred

    best["payload"] = model_payload(model, model_feature_columns, selected_calibration)
    save_model_payload(MODEL_KEYS[name], plant, model, model_feature_columns, selected_calibration)
    audit = {
        "model": name,
        "plant": plant,
        "feature_count": len(model_feature_columns),
        "selected_hyperparameters": best["config"],
        "one_step_validation_score": best["screen_score"],
        "rolling_validation_score": selected_rolling_score,
        "rolling_validation_metrics": selected_rolling_metrics,
        "calibration": {
            "shrinkage": selected_calibration.get("shrinkage"),
            "bias_correction_mw": selected_calibration.get("bias_correction_mw"),
            "bin_calibration": bool(selected_calibration.get("bin_calibration")),
            "hourly_residual_correction": bool(selected_calibration.get("hourly_residual_correction")),
            "profile_blend": bool(selected_calibration.get("profile_blend")),
            "level_anchor_blend": selected_calibration.get("level_anchor_blend"),
        },
        "finalists": finalist_scores,
        "testing_used_for_selection_or_calibration": False,
    }
    return best["payload"], audit


# Predicts residual/delta output with saved feature alignment.
def predict_aligned(model_info, X, plant):
    model, model_feature_columns = unpack_model_payload(model_info)
    if not model_feature_columns:
        model_feature_columns = list(X.columns)
    if set(X.columns) != set(model_feature_columns):
        print("WARNING: Feature mismatch detected. Aligning features or retraining model.")
    X_input = align_features(X.copy(), model_feature_columns)
    return model.predict(X_input)


# ============================================================
# FEATURE ENGINEERING AND CHRONOLOGICAL SPLITTING
# ============================================================

# Builds time, lag, rolling, unit-share, outage, hydrologic, and upstream features aligned with RBFNN.
def add_features(df):
    out = df.copy()
    out["datetime"] = pd.to_datetime(out["datetime"])
    hour0 = out["time"].astype(int) - 1
    dt = out["datetime"]
    out["hour_sin"] = np.sin(2 * np.pi * hour0 / 24)
    out["hour_cos"] = np.cos(2 * np.pi * hour0 / 24)
    out["day_sin"] = np.sin(2 * np.pi * dt.dt.dayofweek / 7)
    out["day_cos"] = np.cos(2 * np.pi * dt.dt.dayofweek / 7)
    out["month_sin"] = np.sin(2 * np.pi * dt.dt.month / 12)
    out["month_cos"] = np.cos(2 * np.pi * dt.dt.month / 12)
    out["is_weekend"] = (dt.dt.dayofweek >= 5).astype(int)
    target_dt = dt + pd.Timedelta(hours=1)
    target_hour0 = target_dt.dt.hour
    out["target_hour_sin"] = np.sin(2 * np.pi * target_hour0 / 24)
    out["target_hour_cos"] = np.cos(2 * np.pi * target_hour0 / 24)
    out["target_day_sin"] = np.sin(2 * np.pi * target_dt.dt.dayofweek / 7)
    out["target_day_cos"] = np.cos(2 * np.pi * target_dt.dt.dayofweek / 7)
    out["target_is_weekend"] = (target_dt.dt.dayofweek >= 5).astype(int)

    for plant in PLANTS:
        target = f"total_gen_{plant}"
        out[f"{target}_tplus1"] = out[target].shift(-1)
        out[f"{target}_delta_tplus1"] = out[f"{target}_tplus1"] - out[target]
        out[f"{target}_current"] = out[target]
        for lag in LAGS:
            out[f"{target}_lag{lag}"] = out[target].shift(lag)
        for window in ROLL_WINDOWS:
            roll = out[target].rolling(window, min_periods=max(2, window // 2))
            out[f"{target}_rollmean{window}"] = roll.mean()
            out[f"{target}_rollstd{window}"] = roll.std().fillna(0.0)
            out[f"{target}_rollmin{window}"] = roll.min()
            out[f"{target}_rollmax{window}"] = roll.max()
        out[f"{target}_diff1"] = out[target].diff(1)
        out[f"{target}_diff3"] = out[target].diff(3)
        out[f"{target}_diff24"] = out[target].diff(24)
        out[f"{target}_target_lag24"] = out[target].shift(23)
        out[f"{target}_target_lag168"] = out[target].shift(167)
        unit_gen_cols = [c for c in out.columns if re.fullmatch(fr"gen_{plant}_unit\d+", c)]
        for unit_col in unit_gen_cols:
            out[f"{unit_col}_current"] = out[unit_col]
            total_safe = out[target].replace(0, np.nan)
            out[f"{unit_col}_share_current"] = (out[unit_col] / total_safe).replace([np.inf, -np.inf], np.nan).fillna(0.0)
            for lag in LAGS:
                out[f"{unit_col}_lag{lag}"] = out[unit_col].shift(lag)
                out[f"{unit_col}_share_lag{lag}"] = out[f"{unit_col}_share_current"].shift(lag)
        out_cols = [c for c in out.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
        if out_cols:
            out[f"{plant}_units_running"] = out[out_cols].sum(axis=1)
            out[f"{plant}_plant_available"] = (out[f"{plant}_units_running"] > 0).astype(int)
    for plant, upstream in UPSTREAM_MAP.items():
        if upstream is None:
            continue
        up_target = f"total_gen_{upstream}"
        out[f"{plant}_upstream_current"] = out[up_target]
        for lag in LAGS:
            out[f"{plant}_upstream_gen_lag{lag}"] = out[up_target].shift(lag)
    for col in [c for c in out.columns if c.startswith("tot_agus") or c.startswith("elev_agus") or "outflow" in c or c == "rainfall"]:
        for lag in [1, 3, 6, 12, 24]:
            out[f"{col}_lag{lag}"] = out[col].shift(lag)
    return out


# Selects the benchmark feature columns for one plant.
def feature_cols(df, plant):
    target = f"total_gen_{plant}"
    cols = [
        "hour_sin",
        "hour_cos",
        "day_sin",
        "day_cos",
        "month_sin",
        "month_cos",
        "is_weekend",
        "target_hour_sin",
        "target_hour_cos",
        "target_day_sin",
        "target_day_cos",
        "target_is_weekend",
        f"{target}_current",
    ]
    prefixes = [f"{target}_lag", f"{target}_roll", f"{target}_diff", f"{target}_target_", f"{plant}_upstream_"]
    cols += [c for c in df.columns if any(c.startswith(prefix) for prefix in prefixes)]
    cols += [
        c for c in df.columns
        if re.fullmatch(fr"gen_{plant}_unit\d+_(current|lag\d+|share_current|share_lag\d+)", c)
    ]
    cols += [c for c in df.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    cols += [f"{plant}_units_running", f"{plant}_plant_available"]
    cols += [c for c in df.columns if c.startswith(f"tot_{plant}") or c.startswith(f"elev_{plant}") or c.startswith(f"spill_{plant}")]
    cols += [c for c in df.columns if "_lag" in c and (c.startswith("tot_agus") or c.startswith("elev_agus") or "outflow" in c or c.startswith("rainfall"))]
    selected = [c for c in dict.fromkeys(cols) if c in df.columns]
    assert not any(
        re.fullmatch(r"rainfall|.*outflow.*", c) and "_lag" not in c
        for c in selected
    ), "Contemporaneous rainfall/outflow must not be a feature"
    return selected


# Splits each plant series chronologically into training, validation, and testing sets.
def split(data):
    n = len(data)
    i1 = int(n * 0.70)
    i2 = int(n * 0.85)
    return data.iloc[:i1], data.iloc[i1:i2], data.iloc[i2:]


# ============================================================
# OUTAGE-AWARE UNIT ALLOCATION
# ============================================================

# Extracts the unit number from an outage/status column.
def unit_from_outage_col(col):
    match = re.search(r"unit\d+", col)
    return match.group(0) if match else None


# Reads the latest unit availability from historical cleaned data.
def latest_status(hist_df, plant):
    cols = [c for c in hist_df.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    latest = hist_df.iloc[-1]
    return {c: float(latest[c]) for c in cols}


# Reads planned unit availability for one forecast hour.
def planned_status(planned, idx, plant):
    cols = [c for c in planned.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    return {c: float(planned.loc[idx, c]) for c in cols}


# Compares planned available capacity against the latest observed baseline.
def availability_ratio(plant, baseline, planned):
    if not planned:
        return 1.0
    base_cap = 0.0
    plan_cap = 0.0
    for col, status in baseline.items():
        base_cap += UNIT_CAPACITY[plant].get(unit_from_outage_col(col), 0.0) * status
    for col, status in planned.items():
        plan_cap += UNIT_CAPACITY[plant].get(unit_from_outage_col(col), 0.0) * status
    if plan_cap <= 0:
        return 0.0
    if base_cap <= 0:
        return plan_cap / CAPACITY_MW[plant]
    return plan_cap / base_cap


# Lists unit-generation columns for one plant.
def unit_generation_columns(plant):
    return [f"gen_{plant}_{unit}" for unit in UNIT_CAPACITY[plant]]


# Builds the matching outage/status column name for one plant unit.
def unit_status_col(plant, unit):
    return f"out_{plant}_{unit}"


# Sums available unit capacity under the planned outage status.
def available_capacity(plant, status):
    return sum(
        capacity
        for unit, capacity in UNIT_CAPACITY[plant].items()
        if status.get(unit_status_col(plant, unit), 1.0) > 0
    )


def all_units_available_status(plant):
    return {unit_status_col(plant, unit): 1.0 for unit in UNIT_CAPACITY[plant]}


# Computes historical unit generation shares within a plant total.
def _unit_share_frame(plant, hist):
    total = hist[f"total_gen_{plant}"].replace(0, np.nan)
    shares = hist[unit_generation_columns(plant)].clip(lower=0.0).div(total, axis=0)
    return shares.replace([np.inf, -np.inf], np.nan)


# Learns unit allocation weights from recent, current, and same-hour unit shares.
def learned_unit_weights(plant, hist, status, forecast_hour):
    unit_cols = unit_generation_columns(plant)
    available_cols = [
        f"gen_{plant}_{unit}"
        for unit in UNIT_CAPACITY[plant]
        if status.get(unit_status_col(plant, unit), 1.0) > 0
    ]
    weights = pd.Series(0.0, index=unit_cols, dtype=float)
    if not available_cols:
        return weights

    shares = _unit_share_frame(plant, hist)
    recent = shares.tail(168)
    components = []
    component_weights = []
    current = shares.iloc[-1][available_cols].dropna() if not shares.empty else pd.Series(dtype=float)
    if not current.empty and current.sum() > 0:
        components.append(current / current.sum())
        component_weights.append(0.50)
    median_recent = recent[available_cols].median(skipna=True).dropna()
    if not median_recent.empty and median_recent.sum() > 0:
        components.append(median_recent / median_recent.sum())
        component_weights.append(0.35)
    if "time" in hist.columns:
        hour_series = pd.to_numeric(hist["time"], errors="coerce").astype("Int64")
        same_hour = hist.loc[(hour_series == int(forecast_hour)).fillna(False)]
        same_hour_recent = _unit_share_frame(plant, same_hour.tail(56)) if not same_hour.empty else pd.DataFrame()
        same_hour_median = same_hour_recent[available_cols].median(skipna=True).dropna() if not same_hour_recent.empty else pd.Series(dtype=float)
        if not same_hour_median.empty and same_hour_median.sum() > 0:
            components.append(same_hour_median / same_hour_median.sum())
            component_weights.append(0.15)

    if components:
        combined = sum(w * c.reindex(available_cols).fillna(0.0) for w, c in zip(component_weights, components))
        combined = combined / combined.sum() if combined.sum() > 0 else combined
    else:
        caps = pd.Series({f"gen_{plant}_{unit}": cap for unit, cap in UNIT_CAPACITY[plant].items()})
        combined = caps[available_cols] / caps[available_cols].sum()
    weights.loc[available_cols] = combined.reindex(available_cols).fillna(0.0)
    return weights / weights.sum() if weights.sum() > 0 else weights


# Allocates a plant-level forecast to available units without exceeding capacities.
def allocate_with_unit_caps(plant, total_generation, weights, status):
    unit_values = {col: 0.0 for col in unit_generation_columns(plant)}
    available_cols = [
        f"gen_{plant}_{unit}"
        for unit in UNIT_CAPACITY[plant]
        if status.get(unit_status_col(plant, unit), 1.0) > 0
    ]
    if not available_cols:
        return unit_values

    caps = pd.Series({f"gen_{plant}_{unit}": cap for unit, cap in UNIT_CAPACITY[plant].items()}, dtype=float)
    total = float(np.clip(total_generation, 0.0, caps[available_cols].sum()))
    alloc = pd.Series(0.0, index=available_cols, dtype=float)
    remaining = total
    active = list(available_cols)
    active_weights = weights.reindex(active).fillna(0.0)
    if active_weights.sum() <= 0:
        active_weights = caps[active] / caps[active].sum()
    for _ in range(len(active) + 1):
        if remaining <= 1e-9 or not active:
            break
        active_weights = active_weights / active_weights.sum() if active_weights.sum() > 0 else caps[active] / caps[active].sum()
        proposal = active_weights * remaining
        headroom = caps[active] - alloc[active]
        capped = proposal >= headroom
        alloc.loc[active] += np.minimum(proposal, headroom)
        remaining = total - float(alloc.sum())
        active = [col for col in active if not capped.get(col, False) and caps[col] - alloc[col] > 1e-9]
        active_weights = weights.reindex(active).fillna(0.0)
    for col, value in alloc.items():
        unit_values[col] = float(np.clip(value, 0.0, caps[col]))
    return unit_values


# Distributes one plant forecast into unit-level generation values.
def distribute_to_units(plant, plant_forecast, status, hist, forecast_hour, return_base=False):
    base_status = all_units_available_status(plant)
    weights = learned_unit_weights(plant, hist, base_status, forecast_hour)
    base_values = allocate_with_unit_caps(plant, plant_forecast, weights, base_status)
    unit_values = dict(base_values)
    for unit in UNIT_CAPACITY[plant]:
        if status.get(unit_status_col(plant, unit), 1.0) <= 0:
            unit_values[f"gen_{plant}_{unit}"] = 0.0
    if return_base:
        return unit_values, base_values
    return unit_values


# Enforces outages, unit caps, and plant-total consistency in the forecast table.
def validate_and_fix_unit_forecast(forecast, planned):
    for idx in forecast.index:
        for plant in PLANTS:
            unit_cols = unit_generation_columns(plant)
            p_status = planned_status(planned, idx, plant)
            for unit, capacity in UNIT_CAPACITY[plant].items():
                col = f"gen_{plant}_{unit}"
                if p_status.get(unit_status_col(plant, unit), 1.0) <= 0:
                    forecast.loc[idx, col] = 0.0
                forecast.loc[idx, col] = float(np.clip(forecast.loc[idx, col], 0.0, capacity))
            active = [col for col in unit_cols if forecast.loc[idx, col] > 1e-6]
            active_values = pd.to_numeric(forecast.loc[idx, active], errors="coerce").to_numpy(dtype=float) if active else np.array([])
            if len(active_values) > 1 and np.allclose(active_values, active_values[0], atol=1e-6):
                warnings.warn(
                    f"{plant} forecast row {idx + 1} has identical active unit generation; "
                    "check whether historical unit data contains identical unit behavior.",
                    RuntimeWarning,
                )
            total_col = f"total_gen_{plant}"
            unit_sum = float(forecast.loc[idx, unit_cols].sum())
            if not np.isclose(float(forecast.loc[idx, total_col]), unit_sum, atol=1e-6):
                forecast.loc[idx, total_col] = unit_sum
    forecast[CASCADE_FORECAST_COLUMN] = forecast[TOTAL_FORECAST_COLUMNS].sum(axis=1)
    return forecast


# ============================================================
# FORECAST INPUT AND OUTPUT FORMATTING
# ============================================================

# Creates an empty 24-hour forecast table with all plant/unit outputs.
def empty_forecast_frame(planned):
    forecast = pd.DataFrame({"Date": planned["Date"].dt.date, "Hour": planned["Hour"].astype(int)})
    for col in UNIT_FORECAST_COLUMNS + TOTAL_FORECAST_COLUMNS:
        forecast[col] = 0.0
    forecast[CASCADE_FORECAST_COLUMN] = 0.0
    return forecast


# Converts 1-24 forecast hours into display labels.
def display_hour(hour):
    hour0 = int(hour) - 1
    return "00:00" if hour0 == 0 else f"{hour0}:00"


# Converts planned outage hour values into internal 1-24 format.
def parse_planned_hour(value):
    if isinstance(value, str):
        text = value.strip()
        match = re.fullmatch(r"(\d{1,2})(?::00)?", text)
        if match:
            hour0 = int(match.group(1))
            if 0 <= hour0 <= 23:
                return hour0 + 1
    return int(pd.to_numeric(value))


# Converts internal forecast column names into readable output headers.
def forecast_display_column(col):
    unit_match = re.fullmatch(r"gen_agus(\d+)_(unit\d+)", col)
    if unit_match:
        plant_num, unit = unit_match.groups()
        return f"Gen_Agus{plant_num}_Unit{unit.replace('unit', '', 1)}_MW"

    total_match = re.fullmatch(r"total_gen_agus(\d+)", col)
    if total_match:
        return f"Total_Gen_Agus{total_match.group(1)}_MW"

    if col == CASCADE_FORECAST_COLUMN:
        return "Total_Cascade_Generation_MW"
    return col


# Formats the final forecast workbook columns and hour labels.
def format_forecast_output(forecast):
    formatted = forecast.copy()
    formatted["Hour"] = formatted["Hour"].apply(display_hour)
    formatted = formatted.rename(columns={col: forecast_display_column(col) for col in formatted.columns})
    return formatted


# Loads and standardizes the 24-hour planned outage input workbook.
def load_planned():
    planned_path = OUTAGES_DIR / "Planned_Outages_Input.xlsx"
    try:
        planned = pd.read_excel(planned_path)
    except zipfile.BadZipFile as exc:
        raise RuntimeError(
            f"Cannot read planned outage workbook: {planned_path}\n"
            "The file is damaged or was not saved as a valid .xlsx workbook. "
            "Close it in Excel, then rerun Cell 1 to regenerate "
            "outputs/outages_planning/Planned_Outages_Input.xlsx before rerunning this cell."
        ) from exc
    planned.columns = [str(c).strip() for c in planned.columns]
    col_map = {c.lower(): c for c in planned.columns}
    missing = [col for col in ["date", "hour"] if col not in col_map]
    if missing:
        raise ValueError(f"Planned outage workbook is missing required column(s): {', '.join(missing)}")
    planned = planned.rename(columns={col_map["date"]: "Date", col_map["hour"]: "Hour"})
    planned["Date"] = pd.to_datetime(planned["Date"])
    planned["Hour"] = planned["Hour"].apply(parse_planned_hour).astype(int)
    for col in [c for c in planned.columns if re.fullmatch(r"out_agus[124567]_unit\d+", c)]:
        planned[col] = np.where(pd.to_numeric(planned[col], errors="coerce").fillna(1) > 0, 1, 0)
    return planned


# ============================================================
# RECURSIVE 24-HOUR BENCHMARK FORECASTING
# ============================================================

# Builds one forecast feature row from the latest historical data.
def feature_row_from_history(hist, plant, date_val, hour_val, model_features=None, hist_feat=None):
    if hist_feat is None:
        hist_feat = add_features(hist.tail(FEATURE_HISTORY_WINDOW).copy())
    x_cols = list(model_features) if model_features else feature_cols(hist_feat, plant)
    row = hist_feat.iloc[-1].to_dict()
    dt = pd.to_datetime(date_val) + pd.Timedelta(hours=int(hour_val) - 1)
    hour0 = int(hour_val) - 1
    row.update({
        "hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "day_sin": np.sin(2 * np.pi * dt.dayofweek / 7),
        "day_cos": np.cos(2 * np.pi * dt.dayofweek / 7),
        "month_sin": np.sin(2 * np.pi * dt.month / 12),
        "month_cos": np.cos(2 * np.pi * dt.month / 12),
        "is_weekend": int(dt.dayofweek >= 5),
        "target_hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "target_hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "target_day_sin": np.sin(2 * np.pi * dt.dayofweek / 7),
        "target_day_cos": np.cos(2 * np.pi * dt.dayofweek / 7),
        "target_is_weekend": int(dt.dayofweek >= 5),
    })
    return pd.DataFrame([{c: row.get(c, 0.0) if pd.notna(row.get(c, 0.0)) else 0.0 for c in x_cols}])


# Generates the recursive 24-hour benchmark forecast with outage-aware unit allocation.
def forecast_benchmark_24h(raw_df, planned, models_by_plant):
    forecast = empty_forecast_frame(planned)
    hist = raw_df.copy()
    baseline_status = {plant: latest_status(hist, plant) for plant in PLANTS}

    for step in range(24):
        date_i = planned.loc[step, "Date"]
        hour_i = int(planned.loc[step, "Hour"])
        hist_feat = add_features(hist.tail(FEATURE_HISTORY_WINDOW).copy())
        new_row = hist.iloc[-1].copy()
        new_row["date"] = pd.to_datetime(date_i)
        new_row["time"] = hour_i
        new_row["datetime"] = pd.to_datetime(date_i) + pd.Timedelta(hours=hour_i - 1)

        for plant in PLANTS:
            target = f"total_gen_{plant}"
            model, model_features = unpack_model_payload(models_by_plant[plant])
            calibration = model_calibration(models_by_plant[plant])
            x_row = feature_row_from_history(hist, plant, date_i, hour_i, model_features, hist_feat=hist_feat)
            if model_features:
                x_row = align_features(x_row, model_features)
            delta = float(model.predict(x_row)[0])
            last_val = float(hist[target].iloc[-1])
            pred = float(benchmark_level_prediction([last_val], [delta], plant, calibration, hours=[hour_i], anchor=[value_at_lag(hist, target, 23)])[0])
            diffs = hist[target].diff().abs().dropna()
            limit = float(max(1.0, diffs.quantile(0.98))) if not diffs.empty else 2.0
            pred = float(np.clip(pred, last_val - limit, last_val + limit))
            pred = float(np.clip(pred, 0.0, CAPACITY_MW[plant] * 1.05))

            p_status = planned_status(planned, step, plant)
            unit_values, recursive_unit_values = distribute_to_units(plant, pred, p_status, hist, hour_i, return_base=True)
            for unit_col, value in unit_values.items():
                forecast.loc[step, unit_col] = value
            for unit_col, value in recursive_unit_values.items():
                new_row[unit_col] = value
            forecast.loc[step, f"total_gen_{plant}"] = sum(unit_values.values())
            new_row[target] = sum(recursive_unit_values.values())
            for out_col, value in all_units_available_status(plant).items():
                new_row[out_col] = value

        hist = pd.concat([hist, pd.DataFrame([new_row])], ignore_index=True)
        forecast.loc[step, CASCADE_FORECAST_COLUMN] = forecast.loc[step, TOTAL_FORECAST_COLUMNS].sum()
    ordered_cols = ["Date", "Hour"] + UNIT_FORECAST_COLUMNS + TOTAL_FORECAST_COLUMNS + [CASCADE_FORECAST_COLUMN]
    return validate_and_fix_unit_forecast(forecast[ordered_cols], planned)


# ============================================================
# LEAKAGE-SAFE ROLLING 24-HOUR DAY-AHEAD BACKTEST
# ============================================================

def valid_24h_dates(frame):
    counts = frame.groupby(pd.to_datetime(frame["datetime"]).dt.date).size()
    return {pd.Timestamp(day) for day, count in counts.items() if count == 24}


def day_ahead_backtest_days(raw_df):
    n = len(raw_df)
    train_end = int(n * 0.70)
    val_end = int(n * 0.85)
    full_dates = valid_24h_dates(raw_df)

    def usable_days(rows):
        days = []
        for day in sorted(valid_24h_dates(rows)):
            if day in full_dates and not raw_df[raw_df["datetime"] < day].empty:
                days.append(day)
        return days

    return {
        "validation": usable_days(raw_df.iloc[train_end:val_end].copy()),
        "testing": usable_days(raw_df.iloc[val_end:].copy()),
    }


def model_selection_days(days):
    days = list(days)
    if len(days) <= SELECTION_DAY_AHEAD_MAX_DAYS:
        return days
    positions = np.linspace(0, len(days) - 1, SELECTION_DAY_AHEAD_MAX_DAYS).round().astype(int)
    return [days[int(pos)] for pos in sorted(set(positions))]


def outage_plan_from_history(history, forecast_day):
    """Build a leakage-safe unit outage plan from the last row before forecast_day."""
    planned = pd.DataFrame({
        "Date": [pd.Timestamp(forecast_day).normalize()] * 24,
        "Hour": list(range(1, 25)),
    })
    if history.empty:
        raise ValueError(f"No historical outage context available before {forecast_day}.")
    if pd.to_datetime(history["datetime"]).max() >= pd.Timestamp(forecast_day):
        raise ValueError("Outage history includes forecast-day rows; this would leak same-day unit outage status.")
    latest = history.iloc[-1]
    for col in [c for c in history.columns if re.fullmatch(r"out_agus[124567]_unit\d+", c)]:
        value = pd.to_numeric(pd.Series([latest.get(col, 1)]), errors="coerce").fillna(1).iloc[0]
        planned[col] = 1 if value > 0 else 0
    return planned


def historical_proxy_outage_plan(actual_day):
    raise RuntimeError("Do not use same-day actual outage status for validation/testing backtests.")


def history_before_forecast_day(raw_df, forecast_day):
    return raw_df[raw_df["datetime"] < pd.Timestamp(forecast_day)].copy().reset_index(drop=True)


def actual_forecast_day(raw_df, forecast_day):
    start = pd.Timestamp(forecast_day)
    end = start + pd.Timedelta(days=1)
    return raw_df[(raw_df["datetime"] >= start) & (raw_df["datetime"] < end)].copy().reset_index(drop=True)


def backtest_feature_row_from_history(hist, plant, date_val, hour_val, model_features):
    hist_feat = add_features(hist.tail(FEATURE_HISTORY_WINDOW).copy())
    return backtest_feature_row_from_dict(hist_feat.iloc[-1].to_dict(), plant, date_val, hour_val, model_features)


def backtest_feature_row_from_dict(row, plant, date_val, hour_val, model_features):
    if not model_features:
        raise ValueError(f"Missing saved feature columns for {plant}; rerun Cell 3 with --train.")
    x_cols = list(model_features)
    dt = pd.to_datetime(date_val) + pd.Timedelta(hours=int(hour_val) - 1)
    hour0 = int(hour_val) - 1
    row.update({
        "hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "day_sin": np.sin(2 * np.pi * dt.dayofweek / 7),
        "day_cos": np.cos(2 * np.pi * dt.dayofweek / 7),
        "month_sin": np.sin(2 * np.pi * dt.month / 12),
        "month_cos": np.cos(2 * np.pi * dt.month / 12),
        "is_weekend": int(dt.dayofweek >= 5),
        "target_hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "target_hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "target_day_sin": np.sin(2 * np.pi * dt.dayofweek / 7),
        "target_day_cos": np.cos(2 * np.pi * dt.dayofweek / 7),
        "target_is_weekend": int(dt.dayofweek >= 5),
    })
    return pd.DataFrame([{c: row.get(c, 0.0) if pd.notna(row.get(c, 0.0)) else 0.0 for c in x_cols}])


def value_at_lag(frame, column, lag, default=0.0):
    if column not in frame.columns or len(frame) <= lag:
        return default
    value = frame[column].iloc[-(lag + 1)]
    return default if pd.isna(value) else float(value)


def latest_value(frame, column, default=0.0):
    if column not in frame.columns or frame.empty:
        return default
    value = frame[column].iloc[-1]
    return default if pd.isna(value) else float(value)


def fast_backtest_feature_row(hist, plant, date_val, hour_val, model_features):
    dt = pd.to_datetime(date_val) + pd.Timedelta(hours=int(hour_val) - 1)
    hour0 = int(hour_val) - 1
    target = f"total_gen_{plant}"
    out_cols = [c for c in hist.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    feature_values = {
        "hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "day_sin": np.sin(2 * np.pi * dt.dayofweek / 7),
        "day_cos": np.cos(2 * np.pi * dt.dayofweek / 7),
        "month_sin": np.sin(2 * np.pi * dt.month / 12),
        "month_cos": np.cos(2 * np.pi * dt.month / 12),
        "is_weekend": int(dt.dayofweek >= 5),
        "target_hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "target_hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "target_day_sin": np.sin(2 * np.pi * dt.dayofweek / 7),
        "target_day_cos": np.cos(2 * np.pi * dt.dayofweek / 7),
        "target_is_weekend": int(dt.dayofweek >= 5),
        f"{target}_current": latest_value(hist, target),
        f"{plant}_units_running": sum(latest_value(hist, c, 1.0) for c in out_cols),
    }
    feature_values[f"{plant}_plant_available"] = int(feature_values[f"{plant}_units_running"] > 0)

    for col in model_features:
        if col in feature_values:
            continue
        if col in hist.columns:
            feature_values[col] = latest_value(hist, col)
        elif m := re.fullmatch(r"(.+)_lag(\d+)", col):
            feature_values[col] = value_at_lag(hist, m.group(1), int(m.group(2)))
        elif m := re.fullmatch(r"(.+)_rollmean(\d+)", col):
            values = pd.to_numeric(hist[m.group(1)].tail(int(m.group(2))), errors="coerce")
            feature_values[col] = float(values.mean()) if values.notna().any() else 0.0
        elif m := re.fullmatch(r"(.+)_rollstd(\d+)", col):
            values = pd.to_numeric(hist[m.group(1)].tail(int(m.group(2))), errors="coerce")
            feature_values[col] = float(values.std()) if values.notna().sum() > 1 else 0.0
        elif m := re.fullmatch(r"(.+)_rollmin(\d+)", col):
            values = pd.to_numeric(hist[m.group(1)].tail(int(m.group(2))), errors="coerce")
            feature_values[col] = float(values.min()) if values.notna().any() else 0.0
        elif m := re.fullmatch(r"(.+)_rollmax(\d+)", col):
            values = pd.to_numeric(hist[m.group(1)].tail(int(m.group(2))), errors="coerce")
            feature_values[col] = float(values.max()) if values.notna().any() else 0.0
        elif m := re.fullmatch(r"(.+)_diff(\d+)", col):
            base, lag = m.group(1), int(m.group(2))
            feature_values[col] = latest_value(hist, base) - value_at_lag(hist, base, lag)
        elif col == f"{target}_target_lag24":
            feature_values[col] = value_at_lag(hist, target, 23)
        elif col == f"{target}_target_lag168":
            feature_values[col] = value_at_lag(hist, target, 167)
        elif m := re.fullmatch(fr"(gen_{plant}_unit\d+)_share_current", col):
            total = latest_value(hist, target)
            feature_values[col] = latest_value(hist, m.group(1)) / total if total else 0.0
        elif m := re.fullmatch(fr"(gen_{plant}_unit\d+)_share_lag(\d+)", col):
            total_lag = value_at_lag(hist, target, int(m.group(2)))
            feature_values[col] = value_at_lag(hist, m.group(1), int(m.group(2))) / total_lag if total_lag else 0.0
        elif col == f"{plant}_upstream_current" and UPSTREAM_MAP.get(plant):
            feature_values[col] = latest_value(hist, f"total_gen_{UPSTREAM_MAP[plant]}")
        elif m := re.fullmatch(fr"{plant}_upstream_gen_lag(\d+)", col):
            upstream = UPSTREAM_MAP.get(plant)
            feature_values[col] = value_at_lag(hist, f"total_gen_{upstream}", int(m.group(1))) if upstream else 0.0
        else:
            feature_values[col] = 0.0

    return pd.DataFrame([{c: feature_values.get(c, 0.0) for c in model_features}])


def forecast_benchmark_24h_backtest(history, planned, models_by_plant):
    forecast = empty_forecast_frame(planned)
    hist = history.copy()
    baseline_status = {plant: latest_status(hist, plant) for plant in PLANTS}

    for step in range(24):
        date_i = planned.loc[step, "Date"]
        hour_i = int(planned.loc[step, "Hour"])
        new_row = hist.iloc[-1].copy()
        new_row["date"] = pd.to_datetime(date_i)
        new_row["time"] = hour_i
        new_row["datetime"] = pd.to_datetime(date_i) + pd.Timedelta(hours=hour_i - 1)
        for plant in PLANTS:
            target = f"total_gen_{plant}"
            model, model_features = unpack_model_payload(models_by_plant[plant])
            calibration = model_calibration(models_by_plant[plant])
            x_row = fast_backtest_feature_row(hist.tail(FEATURE_HISTORY_WINDOW).copy(), plant, date_i, hour_i, model_features)
            if model_features:
                x_row = align_features(x_row, model_features)
            delta = float(model.predict(x_row)[0])
            last_val = float(hist[target].iloc[-1])
            pred = float(benchmark_level_prediction([last_val], [delta], plant, calibration, hours=[hour_i], anchor=[value_at_lag(hist, target, 23)])[0])
            diffs = hist[target].tail(FEATURE_HISTORY_WINDOW).diff().abs().dropna()
            limit = float(max(1.0, diffs.quantile(0.98))) if not diffs.empty else 2.0
            pred = float(np.clip(pred, last_val - limit, last_val + limit))
            pred = float(np.clip(pred, 0.0, CAPACITY_MW[plant] * 1.05))

            p_status = planned_status(planned, step, plant)
            unit_values = distribute_to_units(plant, pred, p_status, hist, hour_i)
            for unit_col, value in unit_values.items():
                forecast.loc[step, unit_col] = value
                new_row[unit_col] = value
            forecast.loc[step, f"total_gen_{plant}"] = sum(unit_values.values())
            new_row[target] = forecast.loc[step, f"total_gen_{plant}"]
            for out_col, value in p_status.items():
                new_row[out_col] = value

        hist = pd.concat([hist, pd.DataFrame([new_row])], ignore_index=True)
        forecast.loc[step, CASCADE_FORECAST_COLUMN] = forecast.loc[step, TOTAL_FORECAST_COLUMNS].sum()

    ordered_cols = ["Date", "Hour"] + UNIT_FORECAST_COLUMNS + TOTAL_FORECAST_COLUMNS + [CASCADE_FORECAST_COLUMN]
    return forecast[ordered_cols]


def rolling_benchmark_plant_predictions(raw_df, forecast_days, model_name, plant, model_info):
    rows = []
    target = f"total_gen_{plant}"
    model, model_features = unpack_model_payload(model_info)
    calibration = model_calibration(model_info)
    for day in forecast_days:
        actual_day = actual_forecast_day(raw_df, day)
        if len(actual_day) != 24:
            continue
        hist = history_before_forecast_day(raw_df, day)
        if hist.empty:
            continue
        planned = outage_plan_from_history(hist, day)
        baseline_status = latest_status(hist, plant)

        for step in range(24):
            date_i = planned.loc[step, "Date"]
            hour_i = int(planned.loc[step, "Hour"])
            dt_val = pd.to_datetime(date_i) + pd.Timedelta(hours=hour_i - 1)
            new_row = hist.iloc[-1].copy()
            new_row["date"] = pd.to_datetime(date_i)
            new_row["time"] = hour_i
            new_row["datetime"] = dt_val

            x_row = fast_backtest_feature_row(hist.tail(FEATURE_HISTORY_WINDOW).copy(), plant, date_i, hour_i, model_features)
            if model_features:
                x_row = align_features(x_row, model_features)
            delta = float(model.predict(x_row)[0])
            last_val = float(hist[target].iloc[-1])
            pred = float(benchmark_level_prediction([last_val], [delta], plant, calibration, hours=[hour_i], anchor=[value_at_lag(hist, target, 23)])[0])
            diffs = hist[target].tail(FEATURE_HISTORY_WINDOW).diff().abs().dropna()
            limit = float(max(1.0, diffs.quantile(0.98))) if not diffs.empty else 2.0
            pred = float(np.clip(pred, last_val - limit, last_val + limit))
            pred = float(np.clip(pred, 0.0, CAPACITY_MW[plant] * 1.05))

            p_status = planned_status(planned, step, plant)
            unit_values = distribute_to_units(plant, pred, p_status, hist, hour_i)
            for unit_col, value in unit_values.items():
                new_row[unit_col] = value
            predicted = float(sum(unit_values.values()))
            new_row[target] = predicted
            for out_col, value in p_status.items():
                new_row[out_col] = value

            rows.append({
                "forecast_date": pd.Timestamp(day).date(),
                "forecast_hour": hour_i,
                "datetime": dt_val,
                "plant": plant,
                "actual_generation": float(actual_day.loc[actual_day["datetime"] == dt_val, target].iloc[0]),
                "predicted_generation": predicted,
                "model": model_name,
                "evaluation_type": DAY_AHEAD_EVALUATION_TYPE,
            })
            hist = pd.concat([hist, pd.DataFrame([new_row])], ignore_index=True)
    return pd.DataFrame(rows)


def rolling_benchmark_candidate_score(raw_df, forecast_days, model_name, plant, model_info):
    predictions = rolling_benchmark_plant_predictions(raw_df, forecast_days, model_name, plant, model_info)
    if predictions.empty:
        return (float("inf"), float("inf"), float("inf")), {}
    mape, mae, rmse, r2 = day_ahead_metric_values(
        predictions["actual_generation"],
        predictions["predicted_generation"],
        plant,
    )
    metrics_out = {
        "operational_mape": mape,
        "mae": mae,
        "rmse": rmse,
        "r2": r2,
        "forecast_days": int(predictions["forecast_date"].nunique()),
        "forecast_hours": int(len(predictions)),
    }
    return candidate_score(metrics_out), metrics_out


def day_ahead_prediction_rows(model_name, forecast_day, actual_day, forecast):
    rows = []
    actual_by_time = actual_day.set_index("datetime")
    for idx in forecast.index:
        dt_val = pd.to_datetime(forecast.loc[idx, "Date"]) + pd.Timedelta(hours=int(forecast.loc[idx, "Hour"]) - 1)
        for plant in PLANTS:
            rows.append({
                "forecast_date": pd.Timestamp(forecast_day).date(),
                "forecast_hour": int(forecast.loc[idx, "Hour"]),
                "datetime": dt_val,
                "plant": plant,
                "actual_generation": float(actual_by_time.loc[dt_val, f"total_gen_{plant}"]),
                "predicted_generation": float(forecast.loc[idx, f"total_gen_{plant}"]),
                "model": model_name,
                "evaluation_type": DAY_AHEAD_EVALUATION_TYPE,
            })
    return rows


def day_ahead_metric_values(y_true, y_pred, plant):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    mask = np.abs(y_true) >= operational_threshold(plant)
    mape = np.nan if not mask.any() else float(np.mean(np.abs((y_true[mask] - y_pred[mask]) / np.abs(y_true[mask]))) * 100.0)
    mae = mean_absolute_error(y_true, y_pred)
    rmse = math.sqrt(mean_squared_error(y_true, y_pred))
    r2 = r2_score(y_true, y_pred) if len(y_true) > 1 else np.nan
    return mape, mae, rmse, r2


def day_ahead_mape_counts(y_true, plant):
    y_true = np.asarray(y_true, dtype=float)
    mask = np.abs(y_true) >= operational_threshold(plant)
    included = int(mask.sum())
    excluded = int(len(mask) - included)
    excluded_fraction = float(excluded / len(mask)) if len(mask) else np.nan
    return included, excluded, excluded_fraction


def day_ahead_metrics_frame(predictions, model_name):
    rows = []
    df = pd.DataFrame(predictions)
    for plant, group in df.groupby("plant", sort=False):
        mape, mae, rmse, r2 = day_ahead_metric_values(group["actual_generation"], group["predicted_generation"], plant)
        included, excluded, excluded_fraction = day_ahead_mape_counts(group["actual_generation"], plant)
        rows.append({
            "model": model_name,
            "plant": plant,
            "number_of_forecast_days": int(group["forecast_date"].nunique()),
            "number_of_forecast_hours": int(len(group)),
            "mape_rows_included": included,
            "mape_rows_excluded": excluded,
            "mape_rows_excluded_fraction": excluded_fraction,
            "low_load_regime": bool(excluded_fraction > 0.40),
            "operational_mape": mape,
            "mae": mae,
            "rmse": rmse,
            "r2": r2,
            "outage_input_type": OUTAGE_INPUT_TYPE,
            "hydrologic_input_type": HYDROLOGIC_INPUT_TYPE,
            "evaluation_type": DAY_AHEAD_EVALUATION_TYPE,
        })
    return pd.DataFrame(rows)


def day_ahead_summary_frame(validation_metrics, testing_metrics, feature_counts):
    validation = validation_metrics.rename(
        columns={
            "operational_mape": "val_operational_mape",
            "mae": "val_mae",
            "rmse": "val_rmse",
            "r2": "val_r2",
        }
    )
    testing = testing_metrics.rename(
        columns={
            "operational_mape": "test_operational_mape",
            "mae": "test_mae",
            "rmse": "test_rmse",
            "r2": "test_r2",
        }
    )
    summary = validation[["model", "plant", "val_operational_mape", "val_mae", "val_rmse", "val_r2"]].merge(
        testing[["plant", "test_operational_mape", "test_mae", "test_rmse", "test_r2"]],
        on="plant",
        how="inner",
    )
    summary["feature_count"] = summary["plant"].map(feature_counts).astype(int)
    return summary[METRICS_COLUMNS]


def testing_prediction_export(predictions):
    out = predictions.copy()
    out["Date"] = pd.to_datetime(out["datetime"]).dt.date
    out["Hour"] = pd.to_numeric(out["forecast_hour"]).astype(int)
    return out[TESTING_PREDICTION_COLUMNS]


def update_day_ahead_backtest_metadata(raw_df, split_days):
    path = DAY_AHEAD_BACKTEST_DIR / "day_ahead_backtest_metadata.json"
    existing = {}
    if path.exists():
        existing = json.loads(path.read_text())
    existing.update({
        "evaluation_type": DAY_AHEAD_EVALUATION_TYPE,
        "models": sorted(set(existing.get("models", []) + ["Random Forest", "XGBoost"])),
        "validation_forecast_days": len(split_days["validation"]),
        "testing_forecast_days": len(split_days["testing"]),
        "data_start": str(raw_df["datetime"].min()),
        "data_end": str(raw_df["datetime"].max()),
        "outage_input_type": OUTAGE_INPUT_TYPE,
        "outage_input_note": OUTAGE_INPUT_NOTE,
        "unit_outage_status_rule": "For validation/testing, each 24-hour forecast day receives only the last-known out_agus*_unit* values from before 00:00 of that day.",
        "same_day_actual_unit_outage_used": False,
        "hydrologic_input_type": HYDROLOGIC_INPUT_TYPE,
        "hydrologic_input_note": HYDROLOGIC_INPUT_NOTE,
        "benchmark_model_fitting_control": "Random Forest and XGBoost estimators are fitted on the training split only; saved models are reused for validation/testing backtests.",
        "validation_testing_separation": "Validation may be used for selection if applicable. Testing data are reserved for final evaluation only.",
        "thesis_ready_note": DAY_AHEAD_THESIS_NOTE,
    })
    path.write_text(json.dumps(existing, indent=2))
    print("Saved:", path)


def save_benchmark_day_ahead_backtests(raw_df, forecast_models, force=False):
    print("Running leakage-safe rolling 24-hour day-ahead backtest for Random Forest and XGBoost...")
    print("Using only historical information available before each forecast day.")
    print("Forecast-day actual generation values are used only for scoring, not as inputs.")
    split_days = day_ahead_backtest_days(raw_df)

    for model_name, models_by_plant in forecast_models.items():
        model_key = model_name.lower().replace(" ", "_")
        feature_counts = {
            plant: len(unpack_model_payload(models_by_plant[plant])[1])
            for plant in PLANTS
        }
        split_metrics = {}
        split_predictions = {}
        for split_name, days in split_days.items():
            pred_path = DAY_AHEAD_BACKTEST_DIR / f"{model_key}_{split_name}_day_ahead_predictions.xlsx"
            metrics_path = DAY_AHEAD_BACKTEST_DIR / f"{model_key}_{split_name}_day_ahead_metrics.xlsx"
            if pred_path.exists() and metrics_path.exists() and not force:
                print(f"Skipping existing {model_name} {split_name} day-ahead backtest outputs.")
                pred_df = pd.read_excel(pred_path)
                metrics_df = pd.read_excel(metrics_path)
            else:
                rows = []
                print(f"{model_name} {split_name} rolling day-ahead days: {len(days)}")
                for day in days:
                    actual_day = actual_forecast_day(raw_df, day)
                    if len(actual_day) != 24:
                        continue
                    hist = history_before_forecast_day(raw_df, day)
                    planned = outage_plan_from_history(hist, day)
                    forecast = forecast_benchmark_24h_backtest(hist, planned, models_by_plant)
                    rows.extend(day_ahead_prediction_rows(model_name, day, actual_day, forecast))
                pred_df = pd.DataFrame(rows)
                metrics_df = day_ahead_metrics_frame(rows, model_name) if rows else pd.DataFrame()
                pred_df.to_excel(pred_path, index=False)
                metrics_df.to_excel(metrics_path, index=False)
                format_excel(pred_path)
                format_excel(metrics_path)
                print("Saved:", pred_path)
                print("Saved:", metrics_path)
            split_predictions[split_name] = pred_df
            split_metrics[split_name] = metrics_df

        if {"validation", "testing"}.issubset(split_metrics):
            metrics_path = META_DIR / f"{model_key}_validation_testing_metrics.xlsx"
            summary = day_ahead_summary_frame(split_metrics["validation"], split_metrics["testing"], feature_counts)
            summary.to_excel(metrics_path, index=False)
            format_excel(metrics_path)
            print("Saved:", metrics_path)

        if "testing" in split_predictions and not split_predictions["testing"].empty:
            predictions_path = META_DIR / f"{model_key}_testing_predictions.xlsx"
            testing_prediction_export(split_predictions["testing"]).to_excel(predictions_path, index=False)
            format_excel(predictions_path)
            print("Saved:", predictions_path)

    update_day_ahead_backtest_metadata(raw_df, split_days)
    print("Saved benchmark rolling day-ahead backtest predictions and metrics.")


# Saves one benchmark forecast to Excel and CSV.
def save_forecast_outputs(forecast, model_key, safe_name):
    out_dir = BENCHMARK_OUTPUT_DIRS[model_key]
    xlsx_path = out_dir / f"Day_Ahead_24H_{safe_name}.xlsx"
    csv_path = out_dir / f"Day_Ahead_24H_{safe_name}.csv"
    output_forecast = format_forecast_output(forecast)
    export_forecast = prepare_forecast_export(output_forecast)
    try:
        export_forecast.to_excel(xlsx_path, index=False)
        format_excel(xlsx_path)
    except PermissionError:
        fallback_xlsx = out_dir / f"Day_Ahead_24H_{safe_name}_regenerated.xlsx"
        export_forecast.to_excel(fallback_xlsx, index=False)
        format_excel(fallback_xlsx)
        print(f"Workbook locked, saved fallback: {fallback_xlsx}")
    export_forecast.to_csv(csv_path, index=False)
    print("Saved:", xlsx_path)
    print("Saved:", csv_path)


# ============================================================
# WORKFLOW ENTRY POINTS
# ============================================================

# Loads cleaned data and planned outages required by benchmark workflows.
def load_latest_inputs():
    clean_path = CLEANED_DATA_DIR / "cleaned_hourly_data.parquet"
    if not clean_path.exists():
        raise FileNotFoundError("Run Thesis Forecasting/scripts/cell1_clean_data.py first.")
    raw_df = pd.read_parquet(clean_path)
    raw_df = add_runtime_compatibility_columns(rebuild_datetime(raw_df))
    planned = load_planned()
    return raw_df, planned


def warn_if_saved_benchmarks_missing_rainfall_features(raw_df):
    if "rainfall" not in raw_df.columns:
        return
    missing = []
    for name, model_key in MODEL_KEYS.items():
        for plant in PLANTS:
            model_path = model_file(model_key, f"{model_key}_{plant}.pkl")
            if not model_path.exists():
                continue
            _, model_features = unpack_model_payload(joblib.load(model_path))
            if model_features and not any(str(col).startswith("rainfall_lag") for col in model_features):
                missing.append(f"{name}:{plant}")
    if missing:
        print(
            "WARNING: Rainfall was added to the cleaned feature set. "
            "Retraining is recommended to ensure benchmark models use the updated hydrologic inputs. "
            f"Saved benchmark models missing rainfall lag features for: {', '.join(missing)}"
        )


# Parses optional benchmark model filters for fast forecast-only dashboard calls.
def selected_benchmark_model_keys_from_args():
    selected = None
    for flag in ["--model", "--models"]:
        if flag in sys.argv:
            idx = sys.argv.index(flag)
            if idx + 1 >= len(sys.argv):
                raise ValueError(f"{flag} requires a value: random_forest, xgboost, or both separated by commas.")
            selected = [item.strip().lower() for item in sys.argv[idx + 1].split(",") if item.strip()]
            break
    if selected is None:
        return None
    valid = set(MODEL_KEYS.values())
    invalid = sorted(set(selected) - valid)
    if invalid:
        raise ValueError(f"Unknown benchmark model(s): {invalid}. Valid choices: {sorted(valid)}")
    return selected


# Loads all saved Random Forest and XGBoost models.
def load_saved_benchmark_models(selected_model_keys=None):
    forecast_models = {"Random Forest": {}, "XGBoost": {}}
    for name, model_key in MODEL_KEYS.items():
        if selected_model_keys is not None and model_key not in selected_model_keys:
            continue
        for plant in PLANTS:
            model_path = model_file(model_key, f"{model_key}_{plant}.pkl")
            if not model_path.exists():
                raise FileNotFoundError(f"Missing saved {name} model: {model_path}. Run with --train first.")
            forecast_models[name][plant] = joblib.load(model_path)
    return {name: models for name, models in forecast_models.items() if models}


# Loads saved benchmark models or trains missing/mismatched models.
def load_or_train_benchmark_models(raw_df):
    warn_if_saved_benchmarks_missing_rainfall_features(raw_df)
    df = add_features(raw_df)
    forecast_models = {"Random Forest": {}, "XGBoost": {}}

    for plant in PLANTS:
        target = f"total_gen_{plant}"
        y_col = f"{target}_delta_tplus1"
        x_cols = feature_cols(df, plant)
        work = df.dropna(subset=[y_col] + x_cols).copy()
        train, _, _ = split(work)
        current_features = train[x_cols].columns.tolist()

        for name, model_key in MODEL_KEYS.items():
            model_path = model_file(model_key, f"{model_key}_{plant}.pkl")
            retrain_model = True
            if model_path.exists():
                saved_payload = joblib.load(model_path)
                model, model_feature_columns = unpack_model_payload(saved_payload)
                if model_feature_columns and set(current_features) == set(model_feature_columns):
                    forecast_models[name][plant] = model_payload(
                        model,
                        model_feature_columns,
                        model_calibration(saved_payload),
                    )
                    retrain_model = False
                    print(f"Loaded {name} model for {plant}")
                else:
                    print("WARNING: Feature mismatch detected. Aligning features or retraining model.")
            else:
                print(f"Missing saved {name} model for {plant}; training a new model.")

            if retrain_model:
                forecast_models[name][plant], _ = train_benchmark_model(name, plant, train, x_cols, y_col)

    return forecast_models


# Recomputes validation/testing metrics and testing predictions from saved models.
def save_saved_model_metrics_and_predictions(raw_df, forecast_models):
    df = add_features(raw_df)
    rows = []
    prediction_rows = {"Random Forest": [], "XGBoost": []}
    fairness_rows = []

    for plant in PLANTS:
        target = f"total_gen_{plant}"
        y_col = f"{target}_delta_tplus1"
        x_cols = feature_cols(df, plant)
        work = df.dropna(subset=[y_col] + x_cols).copy()
        _, val, test = split(work)

        for name, models_by_plant in forecast_models.items():
            model_info = models_by_plant[plant]
            _, model_feature_columns = unpack_model_payload(model_info)
            feature_count = len(model_feature_columns or x_cols)
            leakage_check = leakage_feature_audit(model_feature_columns or x_cols)
            val_delta_pred = predict_aligned(model_info, val[x_cols].copy(), plant)
            test_delta_pred = predict_aligned(model_info, test[x_cols].copy(), plant)
            calibration = model_calibration(model_info)
            val_pred = benchmark_level_prediction(
                val[target].values.astype(float),
                val_delta_pred,
                plant,
                calibration,
                hours=target_hours_from_rows(val),
                anchor=same_hour_target_anchor(val, plant),
            )
            test_pred = benchmark_level_prediction(
                test[target].values.astype(float),
                test_delta_pred,
                plant,
                calibration,
                hours=target_hours_from_rows(test),
                anchor=same_hour_target_anchor(test, plant),
            )
            val_actual = val[f"{target}_tplus1"].values.astype(float)
            test_actual = test[f"{target}_tplus1"].values.astype(float)
            val_m = metrics(val_actual, val_pred, plant)
            test_m = metrics(test_actual, test_pred, plant)
            rows.append({
                "model": name,
                "plant": plant,
                "feature_count": feature_count,
                "val_operational_mape": val_m["operational_mape"],
                "val_mae": val_m["mae"],
                "val_rmse": val_m["rmse"],
                "val_r2": val_m["r2"],
                "test_operational_mape": test_m["operational_mape"],
                "test_mae": test_m["mae"],
                "test_rmse": test_m["rmse"],
                "test_r2": test_m["r2"],
            })
            fairness_rows.append({
                "model": name,
                "plant": plant,
                "feature count": feature_count,
                "target type": "residual/delta target reconstructed to generation level",
                "split strategy": "chronological 70/15/15",
                "operational MAPE definition": "same as RBFNN: excludes near-zero actual generation below plant operational threshold",
                "feature set alignment": "aligned with RBFNN feature families: time, lag, rolling, unit-share, outage, hydrologic context, and upstream cascade features",
                "leakage check": leakage_check,
                "benchmark role": "comparison benchmark only; RBFNN remains the primary model",
            })

            test_datetimes = pd.to_datetime(test["datetime"]) + pd.Timedelta(hours=1)
            for dt_val, actual, predicted in zip(test_datetimes, test_actual, test_pred):
                prediction_rows[name].append({
                    "Date": dt_val.date(),
                    "Hour": int(dt_val.hour) + 1,
                    "datetime": dt_val,
                    "plant": plant,
                    "actual_generation": float(actual),
                    "predicted_generation": float(predicted),
                    "model": name,
                })

    result = pd.DataFrame(rows)[METRICS_COLUMNS]
    for name in ["Random Forest", "XGBoost"]:
        model_key = name.lower().replace(" ", "_")
        metrics_path = META_DIR / f"{model_key}_validation_testing_metrics.xlsx"
        predictions_path = META_DIR / f"{model_key}_testing_predictions.xlsx"
        result[result["model"] == name].to_excel(metrics_path, index=False)
        pd.DataFrame(prediction_rows[name])[TESTING_PREDICTION_COLUMNS].to_excel(predictions_path, index=False)
        format_excel(metrics_path)
        format_excel(predictions_path)
        print("Saved:", metrics_path)
        print("Saved:", predictions_path)
    save_benchmark_fairness_audit(fairness_rows)


# Uses saved benchmark models to generate only operational 24-hour forecasts.
def run_fast_forecast_only():
    raw_df, planned = load_latest_inputs()
    forecast_models = load_saved_benchmark_models(selected_benchmark_model_keys_from_args())
    for name, models_by_plant in forecast_models.items():
        forecast = forecast_benchmark_24h(raw_df, planned, models_by_plant)
        model_key = name.lower().replace(" ", "_")
        safe_name = name.upper().replace(" ", "_")
        save_forecast_outputs(forecast, model_key, safe_name)
    print("Benchmark fast forecast-only mode complete")
    print("Saved benchmark forecasts:", OUT_DIR)


# Uses saved or retrained benchmark models to evaluate metrics and forecast 24 hours.
def run_forecast_only():
    raw_df, planned = load_latest_inputs()
    forecast_models = load_or_train_benchmark_models(raw_df)
    save_saved_model_metrics_and_predictions(raw_df, forecast_models)
    for name, models_by_plant in forecast_models.items():
        forecast = forecast_benchmark_24h(raw_df, planned, models_by_plant)
        model_key = name.lower().replace(" ", "_")
        safe_name = name.upper().replace(" ", "_")
        save_forecast_outputs(forecast, model_key, safe_name)
    print("Benchmark forecasts complete")
    print("Saved benchmark forecasts:", OUT_DIR)
    if "--day-ahead-backtest" in sys.argv:
        save_benchmark_day_ahead_backtests(raw_df, forecast_models)


def run_day_ahead_backtest_only():
    raw_df, _ = load_latest_inputs()
    forecast_models = load_or_train_benchmark_models(raw_df)
    save_loaded_benchmark_selection_audit(forecast_models)
    save_benchmark_day_ahead_backtests(raw_df, forecast_models, force=True)


def save_loaded_benchmark_selection_audit(forecast_models=None):
    if forecast_models is None:
        forecast_models = load_saved_benchmark_models()
    metric_frames = {}
    for name, model_key in MODEL_KEYS.items():
        metrics_path = META_DIR / f"{model_key}_validation_testing_metrics.xlsx"
        if metrics_path.exists():
            metric_frames[name] = pd.read_excel(metrics_path)
    rows = []
    for name, models_by_plant in forecast_models.items():
        for plant, model_info in models_by_plant.items():
            model, model_features = unpack_model_payload(model_info)
            params = model.get_params() if hasattr(model, "get_params") else {}
            selected = {
                key: params.get(key)
                for key in [
                    "n_estimators",
                    "max_depth",
                    "min_samples_leaf",
                    "max_features",
                    "learning_rate",
                    "subsample",
                    "colsample_bytree",
                    "reg_lambda",
                ]
                if key in params
            }
            metrics_row = {}
            metrics_df = metric_frames.get(name)
            if metrics_df is not None:
                match = metrics_df[metrics_df["plant"] == plant]
                if not match.empty:
                    metrics_row = match.iloc[0].to_dict()
            rows.append({
                "model": name,
                "plant": plant,
                "feature_count": len(model_features),
                "selected_hyperparameters": selected,
                "validation_score_used_for_selection": {
                    "operational_mape": metrics_row.get("val_operational_mape"),
                    "rmse": metrics_row.get("val_rmse"),
                    "r2": metrics_row.get("val_r2"),
                },
                "testing_score_after_final_evaluation": {
                    "operational_mape": metrics_row.get("test_operational_mape"),
                    "mae": metrics_row.get("test_mae"),
                    "rmse": metrics_row.get("test_rmse"),
                    "r2": metrics_row.get("test_r2"),
                },
                "selection_source": "saved tuned benchmark model",
                "testing_used_for_selection_or_calibration": False,
            })
    update_model_selection_audit("benchmarks", rows)


# Trains benchmark models, saves metrics/artifacts, and generates forecasts.
def run_training_and_forecast():
    raw_df, planned = load_latest_inputs()
    selected_plants = [arg.lower() for arg in sys.argv[sys.argv.index("--train") + 1:] if not arg.startswith("--")]
    if selected_plants:
        invalid = sorted(set(selected_plants) - set(PLANTS))
        if invalid:
            raise ValueError(f"Unknown plant(s): {invalid}. Valid plants: {PLANTS}")
    # --- Feature preparation shared by Random Forest and XGBoost ---
    df = add_features(raw_df)
    rows = []
    testing_prediction_rows = {"Random Forest": [], "XGBoost": []}
    forecast_models = {"Random Forest": {}, "XGBoost": {}}
    fairness_rows = []
    selection_audit_rows = []
    all_validation_forecast_days = day_ahead_backtest_days(raw_df)["validation"]
    validation_forecast_days = model_selection_days(all_validation_forecast_days)
    print(
        f"Benchmark model selection day-ahead sample: "
        f"{len(validation_forecast_days)} of {len(all_validation_forecast_days)} validation days"
    )

    for plant in PLANTS:
        if selected_plants and plant not in selected_plants:
            for name, model_key in MODEL_KEYS.items():
                model_path = model_file(model_key, f"{model_key}_{plant}.pkl")
                if model_path.exists():
                    forecast_models[name][plant] = joblib.load(model_path)
            continue
        target = f"total_gen_{plant}"
        y_col = f"{target}_delta_tplus1"
        x_cols = feature_cols(df, plant)
        work = df.dropna(subset=[y_col] + x_cols).copy()
        train, val, test = split(work)

        for name in ["Random Forest", "XGBoost"]:
            # --- Model training and validation/testing evaluation ---
            model_info, selection_audit = train_benchmark_model(
                name,
                plant,
                train,
                x_cols,
                y_col,
                val=val,
                raw_df=raw_df,
                validation_forecast_days=validation_forecast_days,
            )
            if selection_audit:
                selection_audit_rows.append(selection_audit)
            _, model_feature_columns = unpack_model_payload(model_info)
            leakage_check = leakage_feature_audit(model_feature_columns)
            val_delta_pred = predict_aligned(model_info, val[x_cols].copy(), plant)
            test_delta_pred = predict_aligned(model_info, test[x_cols].copy(), plant)
            calibration = model_calibration(model_info)
            val_pred = benchmark_level_prediction(
                val[target].values.astype(float),
                val_delta_pred,
                plant,
                calibration,
                hours=target_hours_from_rows(val),
                anchor=same_hour_target_anchor(val, plant),
            )
            test_pred = benchmark_level_prediction(
                test[target].values.astype(float),
                test_delta_pred,
                plant,
                calibration,
                hours=target_hours_from_rows(test),
                anchor=same_hour_target_anchor(test, plant),
            )
            val_actual = val[f"{target}_tplus1"].values.astype(float)
            test_actual = test[f"{target}_tplus1"].values.astype(float)
            val_m = metrics(val_actual, val_pred, plant)
            test_m = metrics(test_actual, test_pred, plant)
            rows.append({
                "model": name,
                "plant": plant,
                "feature_count": len(model_feature_columns),
                "val_operational_mape": val_m["operational_mape"],
                "val_mae": val_m["mae"],
                "val_rmse": val_m["rmse"],
                "val_r2": val_m["r2"],
                "test_operational_mape": test_m["operational_mape"],
                "test_mae": test_m["mae"],
                "test_rmse": test_m["rmse"],
                "test_r2": test_m["r2"],
            })
            fairness_rows.append({
                "model": name,
                "plant": plant,
                "feature count": len(model_feature_columns),
                "target type": "residual/delta target reconstructed to generation level",
                "split strategy": "chronological 70/15/15",
                "operational MAPE definition": "same as RBFNN: excludes near-zero actual generation below plant operational threshold",
                "feature set alignment": "aligned with RBFNN feature families: time, lag, rolling, unit-share, outage, hydrologic context, and upstream cascade features",
                "leakage check": leakage_check,
                "benchmark role": "comparison benchmark only; RBFNN remains the primary model",
            })
            test_datetimes = pd.to_datetime(test["datetime"]) + pd.Timedelta(hours=1)
            for dt_val, actual, predicted in zip(test_datetimes, test_actual, test_pred):
                testing_prediction_rows[name].append({
                    "Date": dt_val.date(),
                    "Hour": int(dt_val.hour) + 1,
                    "datetime": dt_val,
                    "plant": plant,
                    "actual_generation": float(actual),
                    "predicted_generation": float(predicted),
                    "model": name,
                })
            forecast_models[name][plant] = model_info

    result = pd.DataFrame(rows)[METRICS_COLUMNS]
    if selected_plants:
        for model_name, model_key in MODEL_KEYS.items():
            existing_path = META_DIR / f"{model_key}_validation_testing_metrics.xlsx"
            if existing_path.exists():
                existing = pd.read_excel(existing_path)
                existing = existing[~existing["plant"].isin(selected_plants)]
                result = pd.concat([existing, result], ignore_index=True)
        result["plant_order"] = result["plant"].map({plant: idx for idx, plant in enumerate(PLANTS)})
        result["model_order"] = result["model"].map({model: idx for idx, model in enumerate(MODEL_KEYS)})
        result = result.sort_values(["model_order", "plant_order"]).drop(columns=["model_order", "plant_order"]).reset_index(drop=True)
    rf_metrics_path = META_DIR / "random_forest_validation_testing_metrics.xlsx"
    xgb_metrics_path = META_DIR / "xgboost_validation_testing_metrics.xlsx"
    result[result["model"] == "Random Forest"].to_excel(rf_metrics_path, index=False)
    result[result["model"] == "XGBoost"].to_excel(xgb_metrics_path, index=False)
    format_excel(rf_metrics_path)
    format_excel(xgb_metrics_path)
    save_benchmark_fairness_audit(fairness_rows)
    update_model_selection_audit("benchmarks", selection_audit_rows)

    for name, prediction_rows in testing_prediction_rows.items():
        model_key = name.lower().replace(" ", "_")
        predictions_path = META_DIR / f"{model_key}_testing_predictions.xlsx"
        predictions = pd.DataFrame(prediction_rows)[TESTING_PREDICTION_COLUMNS]
        if selected_plants and predictions_path.exists():
            existing_predictions = pd.read_excel(predictions_path)
            existing_predictions = existing_predictions[~existing_predictions["plant"].isin(selected_plants)]
            predictions = pd.concat([existing_predictions, predictions], ignore_index=True)
        predictions.to_excel(predictions_path, index=False)
        format_excel(predictions_path)
        print("Saved:", predictions_path)

    for name, models_by_plant in forecast_models.items():
        forecast = forecast_benchmark_24h(raw_df, planned, models_by_plant)
        model_key = name.lower().replace(" ", "_")
        safe_name = name.upper().replace(" ", "_")
        save_forecast_outputs(forecast, model_key, safe_name)
    save_benchmark_day_ahead_backtests(raw_df, forecast_models)
    print(result.to_string(index=False))
    print("Saved:", rf_metrics_path)
    print("Saved:", xgb_metrics_path)
    print("Saved benchmark forecasts:", OUT_DIR)


# Selects training mode when --train is passed; otherwise runs forecast-only mode.
def main():
    if "--audit-only" in sys.argv:
        save_loaded_benchmark_selection_audit()
    elif "--forecast-only" in sys.argv and "--train" not in sys.argv:
        run_fast_forecast_only()
    elif "--day-ahead-backtest" in sys.argv and "--train" not in sys.argv:
        run_day_ahead_backtest_only()
    elif "--train" in sys.argv:
        run_training_and_forecast()
    else:
        run_forecast_only()


if __name__ == "__main__":
    main()
