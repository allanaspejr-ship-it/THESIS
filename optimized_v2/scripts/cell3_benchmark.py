"""
Optimized Cell 3 benchmarks.

Random Forest and XGBoost are comparison models only. They use the same
chronological split and operational MAPE definition as optimized Cell 2.
"""

import math
import re
import sys
import warnings
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from xgboost import XGBRegressor

warnings.filterwarnings("ignore", category=pd.errors.PerformanceWarning)

PROJECT_DIR = Path(__file__).resolve().parents[2]
OPT_DIR = PROJECT_DIR / "optimized_v2"
LEGACY_OPT_DIR = PROJECT_DIR / "optimized_version"
OUT_DIR = OPT_DIR / "outputs"
CLEANED_DATA_DIR = OUT_DIR / "cleaned_data"
OUTAGES_DIR = OUT_DIR / "outages_planning"
BENCHMARK_DIR = OPT_DIR / "benchmark"
BENCHMARK_OUTPUT_DIRS = {
    "random_forest": BENCHMARK_DIR / "random_forest",
    "xgboost": BENCHMARK_DIR / "xgboost",
}
MODEL_DIRS = {
    "random_forest": OPT_DIR / "models" / "random_forest",
    "xgboost": OPT_DIR / "models" / "xgboost",
}
LEGACY_MODEL_DIRS = {
    "random_forest": LEGACY_OPT_DIR / "models" / "random_forest",
    "xgboost": LEGACY_OPT_DIR / "models" / "xgboost",
}
META_DIR = OPT_DIR / "metadata" / "overall_metrics"

META_DIR.mkdir(parents=True, exist_ok=True)
for folder in [CLEANED_DATA_DIR, OUTAGES_DIR, *BENCHMARK_OUTPUT_DIRS.values(), *MODEL_DIRS.values()]:
    folder.mkdir(parents=True, exist_ok=True)

PLANTS = ["agus1", "agus2", "agus4", "agus5", "agus6", "agus7"]
METRICS_COLUMNS = [
    "model",
    "plant",
    "feature_count",
    "val_operational_mape",
    "val_mae",
    "val_rmse",
    "val_r2",
    "test_operational_mape",
    "test_mae",
    "test_rmse",
    "test_r2",
]
TESTING_PREDICTION_COLUMNS = [
    "Date",
    "Hour",
    "datetime",
    "plant",
    "actual_generation",
    "predicted_generation",
    "model",
]
CAPACITY_MW = {"agus1": 80.0, "agus2": 180.0, "agus4": 158.1, "agus5": 55.0, "agus6": 219.0, "agus7": 54.0}
UNIT_CAPACITY = {
    "agus1": {"unit1": 40.0, "unit2": 40.0},
    "agus2": {"unit1": 60.0, "unit2": 60.0, "unit3": 60.0},
    "agus4": {"unit1": 52.7, "unit2": 52.7, "unit3": 52.7},
    "agus5": {"unit1": 27.5, "unit2": 27.5},
    "agus6": {"unit1": 34.5, "unit2": 34.5, "unit3": 50.0, "unit4": 50.0, "unit5": 50.0},
    "agus7": {"unit1": 27.0, "unit2": 27.0},
}
UNIT_FORECAST_COLUMNS = [
    f"gen_{plant}_{unit}"
    for plant in PLANTS
    for unit in UNIT_CAPACITY[plant]
]
TOTAL_FORECAST_COLUMNS = [f"total_gen_{plant}" for plant in PLANTS]
CASCADE_FORECAST_COLUMN = "total_cascade_generation"
LAGS = [1, 2, 3, 6, 12, 24, 48, 72, 168]
ROLL_WINDOWS = [3, 6, 12, 24, 48, 168]


def operational_threshold(plant):
    return max(1.0, 0.01 * CAPACITY_MW[plant])


def operational_mape(y_true, y_pred, plant):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    mask = np.abs(y_true) >= operational_threshold(plant)
    if not mask.any():
        return np.nan
    return float(np.mean(np.abs((y_true[mask] - y_pred[mask]) / np.abs(y_true[mask]))) * 100.0)


def metrics(y_true, y_pred, plant):
    return {
        "operational_mape": operational_mape(y_true, y_pred, plant),
        "mae": mean_absolute_error(y_true, y_pred),
        "rmse": math.sqrt(mean_squared_error(y_true, y_pred)),
        "r2": r2_score(y_true, y_pred),
    }


def format_excel(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            header = str(column_cells[0].value or "")
            max_len = max([len(header)] + [len(str(cell.value)) for cell in column_cells[1:80] if cell.value is not None])
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 34)
    wb.save(path)


def rebuild_datetime(df):
    out = df.copy()
    out["date"] = pd.to_datetime(out["date"])
    out["time"] = pd.to_numeric(out["time"]).astype(int)
    if "datetime" not in out.columns:
        out["datetime"] = out["date"] + pd.to_timedelta(out["time"] - 1, unit="h")
    else:
        out["datetime"] = pd.to_datetime(out["datetime"])
    return out.sort_values("datetime").reset_index(drop=True)


def add_runtime_compatibility_columns(df):
    # Keep optimized_v2 cleaned files public-facing while supporting models trained on old feature names.
    out = df.copy()
    for plant in PLANTS:
        new_gate = f"tot_{plant}_gate"
        old_gate = f"tot_{plant}"
        if new_gate in out.columns and old_gate not in out.columns:
            out[old_gate] = out[new_gate]
    if "lake_lanao_outflow" in out.columns and "lake_lanao_hourly_outflow_elev_agus7" not in out.columns:
        out["lake_lanao_hourly_outflow_elev_agus7"] = out["lake_lanao_outflow"]
    return out


def model_file(model_key, name):
    current = MODEL_DIRS[model_key] / name
    if current.exists():
        return current
    legacy = LEGACY_MODEL_DIRS[model_key] / name
    if legacy.exists():
        return legacy
    return current


def add_features(df):
    out = df.copy()
    out["datetime"] = pd.to_datetime(out["datetime"])
    hour0 = out["time"].astype(int) - 1
    out["hour_sin"] = np.sin(2 * np.pi * hour0 / 24)
    out["hour_cos"] = np.cos(2 * np.pi * hour0 / 24)
    out["day_of_week"] = out["datetime"].dt.dayofweek
    out["month"] = out["datetime"].dt.month

    for plant in PLANTS:
        target = f"total_gen_{plant}"
        out[f"{target}_tplus1"] = out[target].shift(-1)
        out[f"{target}_current"] = out[target]
        for lag in LAGS:
            out[f"{target}_lag{lag}"] = out[target].shift(lag)
        for window in ROLL_WINDOWS:
            roll = out[target].rolling(window, min_periods=max(2, window // 2))
            out[f"{target}_rollmean{window}"] = roll.mean()
            out[f"{target}_rollstd{window}"] = roll.std().fillna(0.0)
        out_cols = [c for c in out.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
        if out_cols:
            out[f"{plant}_units_running"] = out[out_cols].sum(axis=1)
    return out


def feature_cols(df, plant):
    target = f"total_gen_{plant}"
    cols = ["hour_sin", "hour_cos", "day_of_week", "month", f"{target}_current"]
    cols += [c for c in df.columns if c.startswith(f"{target}_lag") or c.startswith(f"{target}_roll")]
    cols += [c for c in df.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    cols += [f"{plant}_units_running"]
    return [c for c in dict.fromkeys(cols) if c in df.columns]


def split(data):
    n = len(data)
    i1 = int(n * 0.70)
    i2 = int(n * 0.85)
    return data.iloc[:i1], data.iloc[i1:i2], data.iloc[i2:]


def unit_from_outage_col(col):
    match = re.search(r"unit\d+", col)
    return match.group(0) if match else None


def latest_status(hist_df, plant):
    cols = [c for c in hist_df.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    latest = hist_df.iloc[-1]
    return {c: float(latest[c]) for c in cols}


def planned_status(planned, idx, plant):
    cols = [c for c in planned.columns if re.fullmatch(fr"out_{plant}_unit\d+", c)]
    return {c: float(planned.loc[idx, c]) for c in cols}


def availability_ratio(plant, baseline, planned):
    if not planned:
        return 1.0
    base_cap = 0.0
    plan_cap = 0.0
    for col, status in baseline.items():
        base_cap += UNIT_CAPACITY[plant].get(unit_from_outage_col(col), 0.0) * status
    for col, status in planned.items():
        plan_cap += UNIT_CAPACITY[plant].get(unit_from_outage_col(col), 0.0) * status
    if plan_cap <= 0:
        return 0.0
    if base_cap <= 0:
        return plan_cap / CAPACITY_MW[plant]
    return plan_cap / base_cap


def distribute_to_units(plant, plant_forecast, status):
    available_units = []
    for unit, capacity in UNIT_CAPACITY[plant].items():
        out_col = f"out_{plant}_{unit}"
        if status.get(out_col, 1.0) > 0:
            available_units.append((unit, capacity))

    unit_values = {f"gen_{plant}_{unit}": 0.0 for unit in UNIT_CAPACITY[plant]}
    if not available_units:
        return unit_values

    available_capacity = sum(capacity for _, capacity in available_units)
    for unit, capacity in available_units:
        unit_values[f"gen_{plant}_{unit}"] = float(plant_forecast) * capacity / available_capacity
    return unit_values


def empty_forecast_frame(planned):
    forecast = pd.DataFrame({"Date": planned["Date"].dt.date, "Hour": planned["Hour"].astype(int)})
    for col in UNIT_FORECAST_COLUMNS + TOTAL_FORECAST_COLUMNS:
        forecast[col] = 0.0
    forecast[CASCADE_FORECAST_COLUMN] = 0.0
    return forecast


def display_hour(hour):
    hour0 = int(hour) - 1
    return "00:00" if hour0 == 0 else f"{hour0}:00"


def parse_planned_hour(value):
    if isinstance(value, str):
        text = value.strip()
        match = re.fullmatch(r"(\d{1,2})(?::00)?", text)
        if match:
            hour0 = int(match.group(1))
            if 0 <= hour0 <= 23:
                return hour0 + 1
    return int(pd.to_numeric(value))


def forecast_display_column(col):
    unit_match = re.fullmatch(r"gen_agus(\d+)_(unit\d+)", col)
    if unit_match:
        plant_num, unit = unit_match.groups()
        return f"Gen_Agus{plant_num}_Unit{unit.replace('unit', '', 1)}_MW"

    total_match = re.fullmatch(r"total_gen_agus(\d+)", col)
    if total_match:
        return f"Total_Gen_Agus{total_match.group(1)}_MW"

    if col == CASCADE_FORECAST_COLUMN:
        return "Total_Cascade_Generation_MW"
    return col


def format_forecast_output(forecast):
    formatted = forecast.copy()
    formatted["Hour"] = formatted["Hour"].apply(display_hour)
    formatted = formatted.rename(columns={col: forecast_display_column(col) for col in formatted.columns})
    return formatted


def load_planned():
    planned_path = OUTAGES_DIR / "Planned_Outages_Input.xlsx"
    planned = pd.read_excel(planned_path)
    planned.columns = [str(c).strip() for c in planned.columns]
    col_map = {c.lower(): c for c in planned.columns}
    planned = planned.rename(columns={col_map["date"]: "Date", col_map["hour"]: "Hour"})
    planned["Date"] = pd.to_datetime(planned["Date"])
    planned["Hour"] = planned["Hour"].apply(parse_planned_hour).astype(int)
    for col in [c for c in planned.columns if re.fullmatch(r"out_agus[124567]_unit\d+", c)]:
        planned[col] = np.where(pd.to_numeric(planned[col], errors="coerce").fillna(1) > 0, 1, 0)
    return planned


def feature_row_from_history(hist, plant, date_val, hour_val):
    hist_feat = add_features(hist.copy())
    x_cols = feature_cols(hist_feat, plant)
    row = hist_feat.iloc[-1].to_dict()
    dt = pd.to_datetime(date_val) + pd.Timedelta(hours=int(hour_val) - 1)
    hour0 = int(hour_val) - 1
    row.update({
        "hour_sin": np.sin(2 * np.pi * hour0 / 24),
        "hour_cos": np.cos(2 * np.pi * hour0 / 24),
        "day_of_week": dt.dayofweek,
        "month": dt.month,
    })
    return pd.DataFrame([{c: row.get(c, 0.0) if pd.notna(row.get(c, 0.0)) else 0.0 for c in x_cols}])


def forecast_benchmark_24h(raw_df, planned, models_by_plant):
    forecast = empty_forecast_frame(planned)
    hist = raw_df.copy()
    baseline_status = {plant: latest_status(hist, plant) for plant in PLANTS}

    for step in range(24):
        date_i = planned.loc[step, "Date"]
        hour_i = int(planned.loc[step, "Hour"])
        new_row = hist.iloc[-1].copy()
        new_row["date"] = pd.to_datetime(date_i)
        new_row["time"] = hour_i
        new_row["datetime"] = pd.to_datetime(date_i) + pd.Timedelta(hours=hour_i - 1)

        for plant in PLANTS:
            target = f"total_gen_{plant}"
            x_row = feature_row_from_history(hist, plant, date_i, hour_i)
            pred = float(models_by_plant[plant].predict(x_row)[0])
            last_val = float(hist[target].iloc[-1])
            diffs = hist[target].diff().abs().dropna()
            limit = float(max(1.0, diffs.quantile(0.98))) if not diffs.empty else 2.0
            pred = float(np.clip(pred, last_val - limit, last_val + limit))
            pred = float(np.clip(pred, 0.0, CAPACITY_MW[plant] * 1.05))

            p_status = planned_status(planned, step, plant)
            ratio = availability_ratio(plant, baseline_status[plant], p_status)
            adjusted = float(np.clip(pred * ratio, 0.0, CAPACITY_MW[plant] * 1.05)) if ratio > 0 else 0.0
            unit_values = distribute_to_units(plant, adjusted, p_status)
            for unit_col, value in unit_values.items():
                forecast.loc[step, unit_col] = value
            forecast.loc[step, f"total_gen_{plant}"] = sum(unit_values.values())
            new_row[target] = pred
            for out_col, value in baseline_status[plant].items():
                new_row[out_col] = value

        hist = pd.concat([hist, pd.DataFrame([new_row])], ignore_index=True)
        forecast.loc[step, CASCADE_FORECAST_COLUMN] = forecast.loc[step, TOTAL_FORECAST_COLUMNS].sum()
    ordered_cols = ["Date", "Hour"] + UNIT_FORECAST_COLUMNS + TOTAL_FORECAST_COLUMNS + [CASCADE_FORECAST_COLUMN]
    return forecast[ordered_cols]


def save_forecast_outputs(forecast, model_key, safe_name):
    out_dir = BENCHMARK_OUTPUT_DIRS[model_key]
    xlsx_path = out_dir / f"Day_Ahead_24H_{safe_name}.xlsx"
    csv_path = out_dir / f"Day_Ahead_24H_{safe_name}.csv"
    output_forecast = format_forecast_output(forecast)
    try:
        output_forecast.to_excel(xlsx_path, index=False)
        format_excel(xlsx_path)
    except PermissionError:
        fallback_xlsx = out_dir / f"Day_Ahead_24H_{safe_name}_regenerated.xlsx"
        output_forecast.to_excel(fallback_xlsx, index=False)
        format_excel(fallback_xlsx)
        print(f"Workbook locked, saved fallback: {fallback_xlsx}")
    output_forecast.to_csv(csv_path, index=False)
    print("Saved:", xlsx_path)
    print("Saved:", csv_path)


def load_latest_inputs():
    clean_path = CLEANED_DATA_DIR / "cleaned_hourly_data.parquet"
    if not clean_path.exists():
        raise FileNotFoundError("Run optimized_v2/scripts/cell1_clean_data.py first.")
    raw_df = pd.read_parquet(clean_path)
    raw_df = add_runtime_compatibility_columns(rebuild_datetime(raw_df))
    planned = load_planned()
    return raw_df, planned


def load_saved_benchmark_models():
    forecast_models = {"Random Forest": {}, "XGBoost": {}}
    model_keys = {"Random Forest": "random_forest", "XGBoost": "xgboost"}
    for name, model_key in model_keys.items():
        for plant in PLANTS:
            model_path = MODEL_DIRS[model_key] / f"{model_key}_{plant}.pkl"
            model_path = model_file(model_key, f"{model_key}_{plant}.pkl")
            if not model_path.exists():
                raise FileNotFoundError(f"Missing saved {name} model: {model_path}. Run with --train first.")
            forecast_models[name][plant] = joblib.load(model_path)
    return forecast_models


def save_saved_model_metrics_and_predictions(raw_df, forecast_models):
    df = add_features(raw_df)
    rows = []
    prediction_rows = {"Random Forest": [], "XGBoost": []}

    for plant in PLANTS:
        target = f"total_gen_{plant}"
        y_col = f"{target}_tplus1"
        x_cols = feature_cols(df, plant)
        work = df.dropna(subset=[y_col] + x_cols).copy()
        _, val, test = split(work)

        for name, models_by_plant in forecast_models.items():
            model = models_by_plant[plant]
            val_pred = np.clip(model.predict(val[x_cols]), 0.0, CAPACITY_MW[plant] * 1.05)
            test_pred = np.clip(model.predict(test[x_cols]), 0.0, CAPACITY_MW[plant] * 1.05)
            val_m = metrics(val[y_col], val_pred, plant)
            test_m = metrics(test[y_col], test_pred, plant)
            rows.append({
                "model": name,
                "plant": plant,
                "feature_count": len(x_cols),
                "val_operational_mape": val_m["operational_mape"],
                "val_mae": val_m["mae"],
                "val_rmse": val_m["rmse"],
                "val_r2": val_m["r2"],
                "test_operational_mape": test_m["operational_mape"],
                "test_mae": test_m["mae"],
                "test_rmse": test_m["rmse"],
                "test_r2": test_m["r2"],
            })

            test_datetimes = pd.to_datetime(test["datetime"]) + pd.Timedelta(hours=1)
            for dt_val, actual, predicted in zip(test_datetimes, test[y_col], test_pred):
                prediction_rows[name].append({
                    "Date": dt_val.date(),
                    "Hour": int(dt_val.hour) + 1,
                    "datetime": dt_val,
                    "plant": plant,
                    "actual_generation": float(actual),
                    "predicted_generation": float(predicted),
                    "model": name,
                })

    result = pd.DataFrame(rows)[METRICS_COLUMNS]
    for name in ["Random Forest", "XGBoost"]:
        model_key = name.lower().replace(" ", "_")
        metrics_path = META_DIR / f"{model_key}_validation_testing_metrics.xlsx"
        predictions_path = META_DIR / f"{model_key}_testing_predictions.xlsx"
        result[result["model"] == name].to_excel(metrics_path, index=False)
        pd.DataFrame(prediction_rows[name])[TESTING_PREDICTION_COLUMNS].to_excel(predictions_path, index=False)
        format_excel(metrics_path)
        format_excel(predictions_path)
        print("Saved:", metrics_path)
        print("Saved:", predictions_path)


def run_forecast_only():
    raw_df, planned = load_latest_inputs()
    forecast_models = load_saved_benchmark_models()
    save_saved_model_metrics_and_predictions(raw_df, forecast_models)
    for name, models_by_plant in forecast_models.items():
        forecast = forecast_benchmark_24h(raw_df, planned, models_by_plant)
        model_key = name.lower().replace(" ", "_")
        safe_name = name.upper().replace(" ", "_")
        save_forecast_outputs(forecast, model_key, safe_name)
    print("Benchmark forecasts complete")
    print("Saved benchmark forecasts:", BENCHMARK_DIR)


def run_training_and_forecast():
    raw_df, planned = load_latest_inputs()
    df = add_features(raw_df)
    rows = []
    testing_prediction_rows = {"Random Forest": [], "XGBoost": []}
    forecast_models = {"Random Forest": {}, "XGBoost": {}}

    for plant in PLANTS:
        target = f"total_gen_{plant}"
        y_col = f"{target}_tplus1"
        x_cols = feature_cols(df, plant)
        work = df.dropna(subset=[y_col] + x_cols).copy()
        train, val, test = split(work)

        models = {
            "Random Forest": RandomForestRegressor(n_estimators=400, min_samples_leaf=2, random_state=42, n_jobs=1),
            "XGBoost": XGBRegressor(
                n_estimators=500,
                learning_rate=0.03,
                max_depth=4,
                subsample=0.9,
                colsample_bytree=0.9,
                objective="reg:squarederror",
                random_state=42,
                n_jobs=1,
            ),
        }

        for name, model in models.items():
            print(f"Training {name} benchmark for {plant}")
            model.fit(train[x_cols], train[y_col])
            val_pred = np.clip(model.predict(val[x_cols]), 0.0, CAPACITY_MW[plant] * 1.05)
            test_pred = np.clip(model.predict(test[x_cols]), 0.0, CAPACITY_MW[plant] * 1.05)
            val_m = metrics(val[y_col], val_pred, plant)
            test_m = metrics(test[y_col], test_pred, plant)
            rows.append({
                "model": name,
                "plant": plant,
                "feature_count": len(x_cols),
                "val_operational_mape": val_m["operational_mape"],
                "val_mae": val_m["mae"],
                "val_rmse": val_m["rmse"],
                "val_r2": val_m["r2"],
                "test_operational_mape": test_m["operational_mape"],
                "test_mae": test_m["mae"],
                "test_rmse": test_m["rmse"],
                "test_r2": test_m["r2"],
            })
            test_datetimes = pd.to_datetime(test["datetime"]) + pd.Timedelta(hours=1)
            for dt_val, actual, predicted in zip(test_datetimes, test[y_col], test_pred):
                testing_prediction_rows[name].append({
                    "Date": dt_val.date(),
                    "Hour": int(dt_val.hour) + 1,
                    "datetime": dt_val,
                    "plant": plant,
                    "actual_generation": float(actual),
                    "predicted_generation": float(predicted),
                    "model": name,
                })
            model_key = name.lower().replace(" ", "_")
            joblib.dump(model, MODEL_DIRS[model_key] / f"{model_key}_{plant}.pkl")
            forecast_models[name][plant] = model

    result = pd.DataFrame(rows)[METRICS_COLUMNS]
    rf_metrics_path = META_DIR / "random_forest_validation_testing_metrics.xlsx"
    xgb_metrics_path = META_DIR / "xgboost_validation_testing_metrics.xlsx"
    result[result["model"] == "Random Forest"].to_excel(rf_metrics_path, index=False)
    result[result["model"] == "XGBoost"].to_excel(xgb_metrics_path, index=False)
    format_excel(rf_metrics_path)
    format_excel(xgb_metrics_path)

    for name, prediction_rows in testing_prediction_rows.items():
        model_key = name.lower().replace(" ", "_")
        predictions_path = META_DIR / f"{model_key}_testing_predictions.xlsx"
        predictions = pd.DataFrame(prediction_rows)[TESTING_PREDICTION_COLUMNS]
        predictions.to_excel(predictions_path, index=False)
        format_excel(predictions_path)
        print("Saved:", predictions_path)

    for name, models_by_plant in forecast_models.items():
        forecast = forecast_benchmark_24h(raw_df, planned, models_by_plant)
        model_key = name.lower().replace(" ", "_")
        safe_name = name.upper().replace(" ", "_")
        save_forecast_outputs(forecast, model_key, safe_name)
    print(result.to_string(index=False))
    print("Saved:", rf_metrics_path)
    print("Saved:", xgb_metrics_path)
    print("Saved benchmark forecasts:", BENCHMARK_DIR)


def main():
    if "--train" in sys.argv:
        run_training_and_forecast()
    else:
        run_forecast_only()


if __name__ == "__main__":
    main()
